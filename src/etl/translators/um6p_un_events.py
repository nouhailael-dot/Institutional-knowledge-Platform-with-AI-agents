"""
Translator for UM6P_EVENT_VF.xlsx
→ event + event_scoring + event_sector tables.

The final event file splits events into ONE TAB PER SECTOR (Agriculture, Water,
AI, ...). Each sector tab has the same 29-column layout (field headers on row 4,
data from row 5). A separate "UN Key Events" tab holds event-only records with
no scoring.

Key differences from the old single-sheet v8 file, all handled here:
  * events are spread across 9 sector tabs, not one "Event Mapping" sheet
  * columns were REORDERED ('Official website' moved before 'Location'), so this
    reads by HEADER NAME via header_resolver, never by fixed position
  * the score-column typo 'Geograpghy' is aliased to 'Geography' by the resolver
  * 'Potential Strategic Value' and 'Strategic Positioning' were dropped from the
    source and their DB columns were removed, so they are not written
  * the sector is taken from the TAB NAME (authoritative), not per-row Yes/No

Idempotent: events matched by (name, host_actor_id). Existing events are enriched
(empty fields filled), not duplicated.

    python scripts/run_um6p_events.py            # dry run
    python scripts/run_um6p_events.py --apply
"""
from __future__ import annotations

import csv
import logging
import re
import uuid
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Optional

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from src.db.models import Actor, Event, EventScoring, EventSector, Sector, Source
from src.etl.common.text_cleaning import clean_text, slugify
from src.etl.common.header_resolver import HeaderMap, find_header_row

logger = logging.getLogger(__name__)

# ============================================================
# CONFIGURATION
# ============================================================

SOURCE_FILE = "data/raw/UM6P_EVENT_VF.xlsx"
SHEET_NY_EVENTS = "UN Key Events"
AUDIT_LOG = "data/interim/um6p_dashboard_audit.csv"
FAILURE_LOG = "data/interim/um6p_dashboard_failures.csv"

# One tab per sector; the tab name IS the sector. None = event only, no sector tag.
SECTOR_SHEETS = {
    "Agriculture": "AGRI",
    "Water": "WATER",
    "AI": "AI",
    "Energy Transition": "ENERGY",
    "Mining": "MINING",
    "Speciality Chemicals": "CHEM",
    "Healthcare": "HEALTH",
    "African Diaspora": "DIASPORA",
    "CROSS-THEMATIC EVENTS": "CROSS",
}

# Headers used to locate the header row on each sheet type
SECTOR_HINT = ["Event Name", "Organizer", "TOTAL /18", "TIER"]
NY_HINT = ["Event Name", "UN Category", "Official Website"]

SCORING_METHODOLOGY_VERSION = "vf"


# ============================================================
# PARSE HELPERS
# ============================================================

def parse_yes_no(value: object) -> Optional[bool]:
    if value is None:
        return None
    s = str(value).strip().lower()
    if s in ("yes", "y", "true", "1"):
        return True
    if s in ("no", "n", "false", "0"):
        return False
    return None


def parse_score(value: object) -> Optional[int]:
    if value is None:
        return None
    try:
        n = int(float(str(value).strip()))
        return n if 1 <= n <= 3 else None
    except (ValueError, TypeError):
        return None


def parse_total(value: object) -> Optional[int]:
    if value is None:
        return None
    try:
        n = int(float(str(value).strip()))
        return n if 6 <= n <= 18 else None
    except (ValueError, TypeError):
        return None


def normalize_tier(raw: Optional[str]) -> tuple[Optional[str], Optional[str]]:
    if not raw:
        return (None, None)
    raw_clean = str(raw).strip()
    match = re.search(r"tier\s*(\d)", raw_clean.lower())
    if match:
        return (f"tier_{match.group(1)}", raw_clean)
    return (None, raw_clean)


def is_section_header(row: list) -> bool:
    """Only the first cell is filled (e.g. a 'KEY ANNUAL ...' banner)."""
    filled = [c for c in row if c not in (None, "")]
    return len(filled) == 1 and row and row[0] not in (None, "")


@dataclass
class RunStats:
    events_inserted: int = 0
    events_matched_existing: int = 0
    scoring_inserted: int = 0
    scoring_updated: int = 0
    sector_tags_created: int = 0
    sector_tags_existing: int = 0
    ny_events_inserted: int = 0
    ny_events_skipped: int = 0
    rows_failed: int = 0
    skipped_section_headers: int = 0

    def summary(self) -> str:
        return (
            f"  Sector tabs: inserted={self.events_inserted} "
            f"matched={self.events_matched_existing} "
            f"scoring(+{self.scoring_inserted}/~{self.scoring_updated}) "
            f"sector_tags(+{self.sector_tags_created})\n"
            f"  UN Key Events: inserted={self.ny_events_inserted} "
            f"skipped={self.ny_events_skipped}\n"
            f"  section_headers_skipped={self.skipped_section_headers} "
            f"rows_failed={self.rows_failed}"
        )


# ============================================================
# TRANSLATOR
# ============================================================

class Um6pEventsTranslator:

    def __init__(self, project_root: Path, session: Session, apply: bool = False):
        self.project_root = project_root
        self.session = session
        self.apply = apply
        self.source_path = project_root / SOURCE_FILE
        self.stats = RunStats()
        self.audit_fh = self._open_log(AUDIT_LOG,
            ["sheet", "row", "verdict", "event_name", "notes"])
        self.failure_fh = self._open_log(FAILURE_LOG,
            ["sheet", "row", "exception", "raw_summary"])
        self._event_cache: dict[tuple, Event] = {}
        self._sector_id_cache: dict[str, uuid.UUID] = {}

    def _open_log(self, rel_path: str, header: list[str]):
        path = self.project_root / rel_path
        path.parent.mkdir(parents=True, exist_ok=True)
        fh = open(path, "w", newline="", encoding="utf-8")
        w = csv.writer(fh)
        w.writerow(header)
        return (fh, w)

    def _log(self, log_pair, *cells):
        log_pair[1].writerow(cells)

    def _close_logs(self):
        for fh, _ in (self.audit_fh, self.failure_fh):
            fh.close()

    @staticmethod
    def _find_sheet(wb, name):
        target = name.strip().lower()
        for s in wb.sheetnames:
            if s.strip().lower() == target:
                return wb[s]
        return None

    # ---------------- source + lookups ----------------

    def _get_or_create_source(self) -> Source:
        desc = f"UM6P event mapping ingested from {SOURCE_FILE}"
        existing = self.session.execute(
            select(Source).where(Source.description == desc)).scalar_one_or_none()
        if existing:
            return existing
        s = Source(source_id=uuid.uuid4(), source_type="spreadsheet", url=None,
                   retrieved_at=datetime.now(), description=desc, trust_tier="verified")
        self.session.add(s)
        self.session.flush()
        return s

    def _get_sector_id(self, code: str) -> Optional[uuid.UUID]:
        if code in self._sector_id_cache:
            return self._sector_id_cache[code]
        sector = self.session.execute(
            select(Sector).where(Sector.code == code)).scalar_one_or_none()
        if sector:
            self._sector_id_cache[code] = sector.sector_id
            return sector.sector_id
        return None

    def _find_host_actor(self, organizer_text: Optional[str]) -> Optional[Actor]:
        if not organizer_text:
            return None
        first = re.split(r"[/,]", organizer_text)[0].strip()
        if not first:
            return None
        slug = slugify(first)
        return self.session.execute(
            select(Actor).where(Actor.slug == slug,
                                Actor.merged_into_actor_id.is_(None))
        ).scalar_one_or_none()

    # ---------------- event upsert ----------------

    def _upsert_event(self, name: str, host: Optional[Actor], location: Optional[str],
                      description: Optional[str], date_text: Optional[str],
                      website: Optional[str], frequency: Optional[str],
                      sponsoring_orgs: Optional[str], source: Source,
                      thematic_focus: Optional[str] = None,
                      event_type: str = "Strategic event") -> tuple[Event, str]:
        host_id = host.actor_id if host else None
        key = (name.lower(), str(host_id) if host_id else None)
        if key in self._event_cache:
            return (self._event_cache[key], "cached")

        q = select(Event).where(Event.name == name)
        q = q.where(Event.host_actor_id == host_id) if host_id else q.where(Event.host_actor_id.is_(None))
        existing = self.session.execute(q).scalar_one_or_none()

        if existing:
            existing.event_type = existing.event_type or event_type
            existing.location = existing.location or location
            existing.website = existing.website or website
            existing.recurring_pattern = existing.recurring_pattern or frequency
            existing.sponsoring_orgs = existing.sponsoring_orgs or sponsoring_orgs
            existing.description = existing.description or description
            existing.thematic_focus = existing.thematic_focus or thematic_focus
            self._event_cache[key] = existing
            self.stats.events_matched_existing += 1
            return (existing, "matched_existing")

        ev = Event(event_id=uuid.uuid4(), name=name, event_type=event_type,
                   location=location, website=website, recurring_pattern=frequency,
                   sponsoring_orgs=sponsoring_orgs, description=description,
                   thematic_focus=thematic_focus, host_actor_id=host_id,
                   source_id=source.source_id)
        self.session.add(ev)
        self.session.flush()
        self._event_cache[key] = ev
        self.stats.events_inserted += 1
        return (ev, "inserted")

    def _upsert_scoring(self, event: Event, hm: HeaderMap, row: list, source: Source):
        existing = self.session.execute(
            select(EventScoring).where(EventScoring.event_id == event.event_id)
        ).scalar_one_or_none()

        tier_norm, tier_raw = normalize_tier(clean_text(hm.get(row, "TIER")))

        fields = {
            "longevity_description":             clean_text(hm.get(row, "Event Longevity")),
            "recognition_description":           clean_text(hm.get(row, "Event Recognition Level")),
            "scale_description":                 clean_text(hm.get(row, "Audience Size")),
            "audience_profile_description":      clean_text(hm.get(row, "Speakers Profile")),
            "organizer_credibility_description": clean_text(hm.get(row, "Organizer Credibility")),
            "geographic_reach_description":      clean_text(hm.get(row, "Geographic Reach")),
            "longevity_score":                   parse_score(hm.get(row, "Longevity Score")),
            "recognition_score":                 parse_score(hm.get(row, "Recognition Score")),
            "scale_score":                       parse_score(hm.get(row, "Scale Score")),
            "audience_profile_score":            parse_score(hm.get(row, "Audience Profile Score")),
            "organizer_profile_score":           parse_score(hm.get(row, "Organizer Profile Score")),
            "geographic_reach_score":            parse_score(hm.get(row, "Geography Reach Score")),
            "total_score":                       parse_total(hm.get(row, "TOTAL /18")),
            "tier":                              tier_norm,
            "tier_raw":                          tier_raw,
            "scoring_methodology_version":       SCORING_METHODOLOGY_VERSION,
            "scored_at":                         datetime.now(),
        }
        if existing:
            for k, v in fields.items():
                setattr(existing, k, v)
            self.stats.scoring_updated += 1
        else:
            self.session.add(EventScoring(event_scoring_id=uuid.uuid4(),
                                          event_id=event.event_id,
                                          source_id=source.source_id, **fields))
            self.stats.scoring_inserted += 1

    def _upsert_event_sector(self, event: Event, sector_code: str, source: Source):
        sector_id = self._get_sector_id(sector_code)
        if not sector_id:
            logger.warning("Sector code %s not found in DB; skipping tag", sector_code)
            return
        existing = self.session.execute(
            select(EventSector).where(EventSector.event_id == event.event_id,
                                      EventSector.sector_id == sector_id)
        ).scalar_one_or_none()
        if existing:
            self.stats.sector_tags_existing += 1
            return
        self.session.add(EventSector(event_id=event.event_id, sector_id=sector_id))
        self.stats.sector_tags_created += 1

    # ---------------- sector sheets ----------------

    def _process_em_row(self, sheet_name, row_idx, hm, row, sector_code, source):
        name = clean_text(hm.get(row, "Event Name"))
        if not name:
            return
        organizer_text = clean_text(hm.get(row, "Organizer"))
        host = self._find_host_actor(organizer_text)

        event, verdict = self._upsert_event(
            name=name, host=host,
            location=clean_text(hm.get(row, "Location")),
            description=clean_text(hm.get(row, "Event Description")),
            date_text=clean_text(hm.get(row, "Date")),
            website=clean_text(hm.get(row, "Official website")),
            frequency=clean_text(hm.get(row, "Frequency")),
            sponsoring_orgs=organizer_text,
            thematic_focus=clean_text(hm.get(row, "Thematic Focus")),
            source=source, event_type="Strategic event")

        self._upsert_scoring(event, hm, row, source)
        if sector_code:
            self._upsert_event_sector(event, sector_code, source)

        self._log(self.audit_fh, sheet_name, row_idx, verdict, name,
                  f"host={host.name if host else 'none'}; sector={sector_code or '-'}")

    def _run_sector_sheet(self, ws, hm, data_start, sector_code, source):
        for row_idx, row in enumerate(
                ws.iter_rows(min_row=data_start, values_only=True), start=data_start):
            row = list(row)
            if all(c in (None, "") for c in row):
                continue
            if is_section_header(row):
                self.stats.skipped_section_headers += 1
                continue
            try:
                self._process_em_row(ws.title, row_idx, hm, row, sector_code, source)
            except Exception as e:
                self.stats.rows_failed += 1
                raw = str([str(v)[:30] if v else "" for v in row[:5]])
                self._log(self.failure_fh, ws.title, row_idx, str(e), raw)
                logger.warning("%s row %d failed: %s", ws.title, row_idx, e)

    # ---------------- UN Key Events sheet ----------------

    def _process_ny_row(self, row_idx, hm, row, source):
        if is_section_header(row):
            self.stats.skipped_section_headers += 1
            return
        name = clean_text(hm.get(row, "Event Name"))
        if not name:
            return
        organizer = clean_text(hm.get(row, "Organizer"))
        host = self._find_host_actor(organizer)

        ev, verdict = self._upsert_event(
            name=name, host=host,
            location=clean_text(hm.get(row, "Location")),
            description=clean_text(hm.get(row, "Event Description")),
            date_text=clean_text(hm.get(row, "Date / Frequency")),
            website=clean_text(hm.get(row, "Official Website")),
            frequency=None,
            sponsoring_orgs=organizer,
            thematic_focus=clean_text(hm.get(row, "Thematic Focus")),
            source=source, event_type="UN key event")

        if verdict == "inserted":
            self.stats.ny_events_inserted += 1
        else:
            self.stats.ny_events_skipped += 1
        self._log(self.audit_fh, "un_key_events", row_idx, verdict, name,
                  f"host={host.name if host else 'none'}")

    def _run_ny_events(self, ws, source):
        header_row, data_start = find_header_row(ws, NY_HINT)
        hm = HeaderMap(header_row)
        for row_idx, row in enumerate(
                ws.iter_rows(min_row=data_start, values_only=True), start=data_start):
            row = list(row)
            if all(c in (None, "") for c in row):
                continue
            try:
                self._process_ny_row(row_idx, hm, row, source)
            except Exception as e:
                self.stats.rows_failed += 1
                raw = str([str(v)[:30] if v else "" for v in row[:4]])
                self._log(self.failure_fh, "un_key_events", row_idx, str(e), raw)
                logger.warning("UN Key Events row %d failed: %s", row_idx, e)

    # ---------------- entry point ----------------

    def run(self) -> RunStats:
        if not self.source_path.exists():
            raise FileNotFoundError(f"Source file not found: {self.source_path}")
        wb = openpyxl.load_workbook(self.source_path, read_only=True, data_only=True)
        try:
            source = self._get_or_create_source()
            for sheet_name, sector_code in SECTOR_SHEETS.items():
                ws = self._find_sheet(wb, sheet_name)
                if ws is None:
                    logger.warning("sector sheet '%s' not found; skipping", sheet_name)
                    continue
                header_row, data_start = find_header_row(ws, SECTOR_HINT)
                hm = HeaderMap(header_row)
                self._run_sector_sheet(ws, hm, data_start, sector_code, source)
            self.session.flush()

            ny_ws = self._find_sheet(wb, SHEET_NY_EVENTS)
            if ny_ws is not None:
                self._run_ny_events(ny_ws, source)

            if self.apply:
                self.session.commit()
            else:
                self.session.rollback()
        finally:
            wb.close()
            self._close_logs()
        return self.stats