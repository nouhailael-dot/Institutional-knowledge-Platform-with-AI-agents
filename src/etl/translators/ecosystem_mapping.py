"""
Translator for Ecosystem_Mapping_The_Final_UpV.xlsx → hub table.

This is the canonical translator pattern for the project. Future translators
follow this structure.

Behavior:
- Idempotent: rows are matched by hub.slug; existing rows are updated, not duplicated.
- All operations occur in a single transaction.
- Every row receives one of four verdicts (inserted/updated/skipped/failed)
  and is logged to data/interim/ecosystem_mapping_audit.csv.
- Failed rows do not stop the run, but cause a non-zero exit status.

Usage (from project root, with .venv active):
    python -m scripts.run_ecosystem_mapping
"""

from __future__ import annotations

import csv
import logging
import uuid
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Optional

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from src.db.models import Hub, Sector, Source
from src.etl.common.text_cleaning import (
    clean_text,
    is_section_header,
    parse_sector_from_category,
    slugify,
)


logger = logging.getLogger(__name__)


# ============================================================
# CONFIGURATION (per ADR-006)
# ============================================================

# Path inside the project. Resolved against PROJECT_ROOT by the runner.
SOURCE_FILE = "data/raw/Ecosystem_Mapping_The_Final_UpV.xlsx"
SHEET_NAME = "Ecosystem Mapping"
AUDIT_LOG_PATH = "data/interim/ecosystem_mapping_audit.csv"
FAILURE_LOG_PATH = "data/interim/ecosystem_mapping_failures.csv"
SKIPPED_LOG_PATH = "data/interim/ecosystem_mapping_skipped.csv"

# Columns in the spreadsheet (1-indexed in Excel, but openpyxl gives us 0-indexed lists).
# These are the real columns per ADR-006; columns 17-38 are empty padding to drop.
COL_NAME = 0
COL_CATEGORY = 1
COL_DESCRIPTION = 2
COL_COMPOSITION = 3
COL_LOCATION = 4
COL_WEBSITES = 5
COL_VALUE_PROP = 6
COL_CORE_STRENGTHS = 7
COL_TALENT = 8
COL_RESEARCH = 9
COL_INDUSTRY = 10
COL_INTERNATIONAL = 11
COL_PROGRAMS = 12
COL_CONSTRAINTS = 13
COL_FUNDINGS = 14
COL_CROSS_THEMES = 15

DATA_STARTS_AT_ROW = 2  # row 1 is header

# Known data quality corrections from ADR-006.
# Key is the hub name as it appears in the spreadsheet.
KNOWN_ISSUE_CORRECTIONS: dict[str, dict[str, str]] = {
    "Massachusetts Energy Innovation Hub (Boston – Cambridge)": {
        "category": "Ecosystem Hub (Energy – Climate Innovation)",
    },
    "Nevada Mining & Critical Minerals Hub (Reno – Elko – Northern Nevada)": {
        "category": "Ecosystem Hub (Mining & Critical Minerals)",
    },
    "Utah Mining & Mettallurgy Hub": {
        "category": "Ecosystem Hub (Mining & Metallurgy)",
    },
}


# ============================================================
# RUN STATISTICS
# ============================================================

@dataclass
class RunStats:
    """Summary of the translator's run."""
    inserted: int = 0
    updated: int = 0
    skipped_section_headers: int = 0
    skipped_empty: int = 0
    failed: int = 0
    failure_reasons: list[str] = field(default_factory=list)
    sector_unrecognized: list[str] = field(default_factory=list)

    @property
    def total(self) -> int:
        return (self.inserted + self.updated + self.skipped_section_headers
                + self.skipped_empty + self.failed)

    def summary(self) -> str:
        lines = [
            f"  Inserted:                {self.inserted}",
            f"  Updated:                 {self.updated}",
            f"  Skipped (section header):{self.skipped_section_headers}",
            f"  Skipped (empty):         {self.skipped_empty}",
            f"  Failed:                  {self.failed}",
            f"  Total rows examined:     {self.total}",
        ]
        if self.sector_unrecognized:
            lines.append(f"\n  Sectors not recognized (review needed):")
            for s in self.sector_unrecognized:
                lines.append(f"    - {s}")
        if self.failure_reasons:
            lines.append(f"\n  First few failure reasons:")
            for r in self.failure_reasons[:5]:
                lines.append(f"    - {r}")
        return "\n".join(lines)


# ============================================================
# CORE TRANSLATOR
# ============================================================

class EcosystemMappingTranslator:
    """Reads the Ecosystem Mapping spreadsheet and loads it into the hub table.

    Usage:
        translator = EcosystemMappingTranslator(project_root, session)
        stats = translator.run()
    """

    def __init__(self, project_root: Path, session: Session):
        self.project_root = project_root
        self.session = session
        self.source_path = project_root / SOURCE_FILE
        self.stats = RunStats()

        # Open audit log files
        self.audit_log = self._open_log(AUDIT_LOG_PATH,
                                        ["row", "verdict", "hub_name", "sector_code", "notes"])
        self.failure_log = self._open_log(FAILURE_LOG_PATH,
                                          ["row", "hub_name", "exception", "raw_values"])
        self.skipped_log = self._open_log(SKIPPED_LOG_PATH,
                                          ["row", "reason", "raw_values"])

    def _open_log(self, path_str: str, columns: list[str]):
        path = self.project_root / path_str
        path.parent.mkdir(parents=True, exist_ok=True)
        fh = open(path, "w", newline="", encoding="utf-8")
        writer = csv.writer(fh)
        writer.writerow(columns)
        return fh, writer

    def _log(self, log_pair, *cells):
        _, writer = log_pair
        writer.writerow(cells)

    def _close_logs(self):
        for fh, _ in (self.audit_log, self.failure_log, self.skipped_log):
            fh.close()

    # --------------------------------------------------------
    # SOURCE RECORD
    # --------------------------------------------------------

    def _get_or_create_source(self) -> Source:
        """Return the Source row representing this spreadsheet, creating it if needed."""
        description = f"Ecosystem Mapping spreadsheet ingested from {SOURCE_FILE}"
        existing = self.session.execute(
            select(Source).where(Source.description == description)
        ).scalar_one_or_none()
        if existing is not None:
            return existing

        source = Source(
            source_id=uuid.uuid4(),
            source_type="spreadsheet",
            url=None,
            retrieved_at=datetime.now(),
            description=description,
            trust_tier="verified",
        )
        self.session.add(source)
        self.session.flush()
        return source

    # --------------------------------------------------------
    # SECTOR LOOKUP
    # --------------------------------------------------------

    def _get_sector_by_code(self, code: Optional[str]) -> Optional[Sector]:
        if not code:
            return None
        return self.session.execute(
            select(Sector).where(Sector.code == code)
        ).scalar_one_or_none()

    # --------------------------------------------------------
    # ROW PROCESSING
    # --------------------------------------------------------

    def _process_row(self, row_idx: int, row: tuple, source: Source) -> None:
        """Process a single row. Routes to the correct handler based on row shape."""
        row_list = list(row)
        raw_summary = str([str(v)[:50] if v else "" for v in row_list[:6]])

        # Skip section headers (only column 1 has content)
        if is_section_header(row_list):
            self.stats.skipped_section_headers += 1
            self._log(self.audit_log, row_idx, "skipped_section_header",
                      str(row_list[0])[:80], "", "section divider")
            self._log(self.skipped_log, row_idx, "section_header", raw_summary)
            return

        # Skip fully-empty rows
        if not any(v is not None and str(v).strip() != "" for v in row_list):
            self.stats.skipped_empty += 1
            self._log(self.audit_log, row_idx, "skipped_empty", "", "", "all cells empty")
            self._log(self.skipped_log, row_idx, "empty_row", raw_summary)
            return

        # Process as a hub row
        try:
            verdict, hub_name, sector_code, notes = self._upsert_hub(row_idx, row_list, source)
            self._log(self.audit_log, row_idx, verdict, hub_name, sector_code or "", notes)
        except Exception as e:
            self.stats.failed += 1
            self.stats.failure_reasons.append(f"Row {row_idx}: {e}")
            hub_name = str(row_list[COL_NAME])[:80] if row_list[COL_NAME] else "?"
            self._log(self.audit_log, row_idx, "failed", hub_name, "", str(e))
            self._log(self.failure_log, row_idx, hub_name, str(e), raw_summary)
            logger.warning("Row %d failed: %s", row_idx, e)

    def _upsert_hub(self, row_idx: int, row: list, source: Source) -> tuple[str, str, Optional[str], str]:
        """Create or update a single hub row.

        Returns (verdict, hub_name, sector_code, notes).
        """
        name_raw = row[COL_NAME]
        name = clean_text(name_raw)
        if not name:
            raise ValueError("Hub name is missing")

        # Apply known-issue corrections first
        correction = KNOWN_ISSUE_CORRECTIONS.get(name) or KNOWN_ISSUE_CORRECTIONS.get(name.strip())
        category_raw = row[COL_CATEGORY]
        if correction and "category" in correction:
            category_raw = correction["category"]
            notes_prefix = "applied known-issue correction; "
        else:
            notes_prefix = ""

        category = clean_text(category_raw)
        sector_code = parse_sector_from_category(category)
        if category and not sector_code:
            self.stats.sector_unrecognized.append(f"Row {row_idx}: '{category}'")

        slug = slugify(name)

        # Look for existing hub
        existing = self.session.execute(
            select(Hub).where(Hub.slug == slug)
        ).scalar_one_or_none()

        # Build the field set
        location_raw = clean_text(row[COL_LOCATION])
        location_city, state, country = self._parse_location(location_raw)

        primary_sectors = [sector_code] if sector_code else []

        if existing is None:
            hub = Hub(
                hub_id=uuid.uuid4(),
                name=name,
                slug=slug,
                primary_city=location_city,
                state=state,
                country=country,
                region=None,
                description=clean_text(row[COL_DESCRIPTION]),
                composition_narrative=clean_text(row[COL_COMPOSITION]),
                talent_narrative=clean_text(row[COL_TALENT]),
                research_infrastructure_narrative=clean_text(row[COL_RESEARCH]),
                industry_integration_narrative=clean_text(row[COL_INDUSTRY]),
                international_engagement_narrative=clean_text(row[COL_INTERNATIONAL]),
                programs_and_labs_narrative=clean_text(row[COL_PROGRAMS]),
                upstream_inputs=None,
                downstream_outputs=None,
                integrated_flow_description=None,
                constraints_summary=clean_text(row[COL_CONSTRAINTS]),
                funding_summary=clean_text(row[COL_FUNDINGS]),
                cross_hub_themes=clean_text(row[COL_CROSS_THEMES]),
                primary_sectors=primary_sectors,
                secondary_sectors=[],
                source_id=source.source_id,
                verification_status="phase_1",
                confidence=1.0,
            )
            self.session.add(hub)
            self.stats.inserted += 1
            return ("inserted", name, sector_code, notes_prefix + "new hub")
        else:
            # Update mutable fields. Do not change slug or id.
            existing.name = name
            existing.primary_city = location_city
            existing.state = state
            existing.country = country
            existing.description = clean_text(row[COL_DESCRIPTION])
            existing.composition_narrative = clean_text(row[COL_COMPOSITION])
            existing.talent_narrative = clean_text(row[COL_TALENT])
            existing.research_infrastructure_narrative = clean_text(row[COL_RESEARCH])
            existing.industry_integration_narrative = clean_text(row[COL_INDUSTRY])
            existing.international_engagement_narrative = clean_text(row[COL_INTERNATIONAL])
            existing.programs_and_labs_narrative = clean_text(row[COL_PROGRAMS])
            existing.constraints_summary = clean_text(row[COL_CONSTRAINTS])
            existing.funding_summary = clean_text(row[COL_FUNDINGS])
            existing.cross_hub_themes = clean_text(row[COL_CROSS_THEMES])
            existing.primary_sectors = primary_sectors
            existing.source_id = source.source_id
            self.stats.updated += 1
            return ("updated", name, sector_code, notes_prefix + "existing hub updated")
    @staticmethod
    def _parse_location(loc: Optional[str]) -> tuple[Optional[str], Optional[str], Optional[str]]:
        """Crude location parser. The Excel column is free-text like
        'Davis / Sacramento Valley / Central Valley, California, USA'.
        Returns (city, state, country)."""
        if not loc:
            return (None, None, None)
        parts = [p.strip() for p in loc.split(",")]
        if len(parts) == 1:
            return (parts[0], None, None)
        if len(parts) == 2:
            return (parts[0], parts[1], None)
        return (parts[0], parts[1], parts[2])

    # --------------------------------------------------------
    # MAIN ENTRY
    # --------------------------------------------------------

    def run(self) -> RunStats:
        """Execute the translator. Returns run statistics."""
        if not self.source_path.exists():
            raise FileNotFoundError(f"Source file not found: {self.source_path}")

        logger.info("Opening %s", self.source_path)
        wb = openpyxl.load_workbook(self.source_path, read_only=True, data_only=True)
        try:
            if SHEET_NAME not in wb.sheetnames:
                raise ValueError(
                    f"Sheet '{SHEET_NAME}' not found. Available sheets: {wb.sheetnames}"
                )
            ws = wb[SHEET_NAME]

            source = self._get_or_create_source()

            for row_idx, row in enumerate(
                ws.iter_rows(min_row=DATA_STARTS_AT_ROW, values_only=True), start=DATA_STARTS_AT_ROW
            ):
                # Trim row to the real columns only (ignore the 22 empty padding columns)
                row = row[:16]
                self._process_row(row_idx, row, source)

            self.session.commit()
        finally:
            wb.close()
            self._close_logs()

        return self.stats