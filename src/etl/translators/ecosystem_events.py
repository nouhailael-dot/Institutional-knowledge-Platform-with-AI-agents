"""
Translator for Ecosystem_Mapping_The_Final_UpV.xlsx / Events Name sheet.

Loads the simple event inventory into the event table. This is deterministic
spreadsheet migration: no AI is needed.
"""

from __future__ import annotations

import csv
import logging
import uuid
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Optional

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from src.db.models import Event, Source
from src.etl.common.text_cleaning import clean_text


logger = logging.getLogger(__name__)

SOURCE_FILE = "data/raw/Ecosystem_Mapping_The_Final_UpV.xlsx"
SHEET_NAME = "Events Name "
AUDIT_LOG = "data/interim/ecosystem_events_audit.csv"
FAILURE_LOG = "data/interim/ecosystem_events_failures.csv"

DATA_STARTS_AT_ROW = 2

COL_EVENT_NAME = 0
COL_LOCATION = 1
COL_EVENT_TYPE = 2
COL_WEBSITE = 3


@dataclass
class RunStats:
    inserted: int = 0
    updated: int = 0
    skipped: int = 0
    failed: int = 0

    def summary(self) -> str:
        return (
            f"  Events inserted: {self.inserted}\n"
            f"  Events updated:  {self.updated}\n"
            f"  Events skipped:  {self.skipped}\n"
            f"  Rows failed:     {self.failed}"
        )


class EcosystemEventsTranslator:
    def __init__(self, project_root: Path, session: Session):
        self.project_root = project_root
        self.session = session
        self.source_path = project_root / SOURCE_FILE
        self.stats = RunStats()

        self.audit_fh = self._open_log(
            AUDIT_LOG, ["row", "verdict", "event_name", "notes"]
        )
        self.failure_fh = self._open_log(
            FAILURE_LOG, ["row", "event_name", "exception", "raw_summary"]
        )

    def _open_log(self, rel_path: str, header: list[str]):
        path = self.project_root / rel_path
        path.parent.mkdir(parents=True, exist_ok=True)
        fh = open(path, "w", newline="", encoding="utf-8")
        writer = csv.writer(fh)
        writer.writerow(header)
        return (fh, writer)

    def _log(self, log_pair, *cells):
        log_pair[1].writerow(cells)

    def _close_logs(self):
        for fh, _ in (self.audit_fh, self.failure_fh):
            fh.close()

    def _get_or_create_source(self) -> Source:
        desc = f"Ecosystem events sheet ingested from {SOURCE_FILE}"
        existing = self.session.execute(
            select(Source).where(Source.description == desc)
        ).scalar_one_or_none()
        if existing:
            return existing
        source = Source(
            source_id=uuid.uuid4(),
            source_type="spreadsheet",
            url=None,
            retrieved_at=datetime.now(),
            description=desc,
            trust_tier="verified",
        )
        self.session.add(source)
        self.session.flush()
        return source

    def _upsert_event(self, row_idx: int, row: list, source: Source) -> None:
        name = clean_text(row[COL_EVENT_NAME])
        if not name:
            self.stats.skipped += 1
            return

        location = clean_text(row[COL_LOCATION])
        event_type = clean_text(row[COL_EVENT_TYPE])
        website = clean_text(row[COL_WEBSITE])

        existing = self.session.execute(
            select(Event).where(Event.name == name)
        ).scalar_one_or_none()

        if existing:
            existing.location = location or existing.location
            existing.event_type = event_type or existing.event_type
            existing.website = website or existing.website
            existing.source_id = existing.source_id or source.source_id
            self.stats.updated += 1
            self._log(self.audit_fh, row_idx, "updated", name, "")
            return

        event = Event(
            event_id=uuid.uuid4(),
            name=name,
            event_type=event_type,
            location=location,
            website=website,
            source_id=source.source_id,
        )
        self.session.add(event)
        self.stats.inserted += 1
        self._log(self.audit_fh, row_idx, "inserted", name, "")

    def run(self) -> RunStats:
        if not self.source_path.exists():
            raise FileNotFoundError(f"Source file not found: {self.source_path}")

        wb = openpyxl.load_workbook(self.source_path, read_only=True, data_only=True)
        try:
            if SHEET_NAME not in wb.sheetnames:
                raise ValueError(f"Sheet '{SHEET_NAME}' not found. Sheets: {wb.sheetnames}")

            source = self._get_or_create_source()
            ws = wb[SHEET_NAME]
            for row_idx, row in enumerate(
                ws.iter_rows(min_row=DATA_STARTS_AT_ROW, values_only=True),
                start=DATA_STARTS_AT_ROW,
            ):
                row = list(row[:4])
                while len(row) < 4:
                    row.append(None)
                try:
                    self._upsert_event(row_idx, row, source)
                except Exception as err:
                    self.stats.failed += 1
                    name = clean_text(row[COL_EVENT_NAME]) or ""
                    raw_summary = str([str(v)[:50] if v else "" for v in row])
                    self._log(self.failure_fh, row_idx, name, str(err), raw_summary)
                    logger.warning("Event row %d failed: %s", row_idx, err)

            self.session.commit()
        finally:
            wb.close()
            self._close_logs()

        return self.stats
