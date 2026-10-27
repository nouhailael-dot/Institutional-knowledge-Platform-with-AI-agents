"""Read an uploaded Data Feed file safely: .xlsx or .csv, values only.

Refused before any cell is read:
    * anything over MAX_BYTES, or with an extension other than .xlsx / .csv
    * a workbook that is not a plain zip of spreadsheet parts: macros, links to
      external files, too many parts, or parts that inflate suspiciously
      (a compressed bomb)
Cell by cell:
    * formulas are read as their text, never evaluated; the row check reports
      them (and any text that starts like one) as an error on that row
    * a value longer than MAX_CELL characters refuses the file
The file itself is never stored: only its name, size, SHA-256 and the values
read from it. Headers are matched against the template; the rows come back
as {header key: text}, with their spreadsheet row numbers.
"""

from __future__ import annotations

import csv
import hashlib
import io
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any

from src.datafeed import contract

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000
MAX_COLUMNS = 150
MAX_CELL = 4000
# Workbook container limits (a real template is ~30 KB and ~20 parts).
MAX_PARTS = 200
MAX_UNZIPPED = 60 * 1024 * 1024
MAX_RATIO = 200
FORBIDDEN_PARTS = ("vbaproject.bin", "externallink", "activex", "oleobject", "embeddings/")
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


class UploadRefused(ValueError):
    """The file cannot be read; the message says why, plainly."""


@dataclass
class ParsedFile:
    file_name: str
    extension: str
    size: int
    sha256: str
    headers: list[str]
    rows: list[dict[str, Any]]                 # {"row_number", "values": {header: text}}
    template: dict[str, str] = field(default_factory=dict)   # workbook identity, if any
    notes: list[str] = field(default_factory=list)
    blank_rows: int = 0


def formula_like(value: str) -> bool:
    """A text that a spreadsheet would run as a formula (CSV injection).

    A plain signed number ("-12", "+3.5") is a value, not a formula.
    """
    if not value or value[0] not in FORMULA_START:
        return False
    if value[0] in "+-":
        try:
            float(value.replace(",", ""))
            return False
        except ValueError:
            return True
    return True


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def check_container(data: bytes) -> None:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise UploadRefused("This is not a valid .xlsx workbook.") from exc
    with archive:
        infos = archive.infolist()
        if len(infos) > MAX_PARTS:
            raise UploadRefused("The workbook has too many parts to be a GHUS template.")
        names = [i.filename.replace("\\", "/").casefold() for i in infos]
        if "[content_types].xml" not in names or "xl/workbook.xml" not in names:
            raise UploadRefused("This is not a valid .xlsx workbook.")
        for name in names:
            if name.startswith("/") or ".." in name.split("/"):
                raise UploadRefused("The workbook contains an unsafe path.")
            if any(part in name for part in FORBIDDEN_PARTS):
                raise UploadRefused("Workbooks with macros, embedded objects or links to "
                                    "other files are not accepted. Save it as a plain .xlsx.")
        total = sum(i.file_size for i in infos)
        if total > MAX_UNZIPPED:
            raise UploadRefused("The workbook expands to more than 60 MB.")
        for info in infos:
            if info.compress_size and info.file_size / info.compress_size > MAX_RATIO \
                    and info.file_size > 1024 * 1024:
                raise UploadRefused("The workbook is compressed suspiciously and was refused.")


def _xlsx_rows(data: bytes, record_type: str) -> tuple[list[list[Any]], dict[str, str], list[str]]:
    from openpyxl import load_workbook

    check_container(data)
    notes: list[str] = []
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False,
                           keep_links=False)
    except Exception as exc:
        raise UploadRefused("The workbook could not be read. Save it again as .xlsx.") from exc
    try:
        identity: dict[str, str] = {}
        if "Lists" in wb.sheetnames:
            marker = [_cell_text(c.value) for c in next(wb["Lists"].iter_rows(max_row=1))][:3]
            if len(marker) == 3 and marker[0] == "ghus-datafeed":
                identity = {"record_type": marker[1], "template_version": marker[2]}
        sheet = contract.SHEET_NAMES[record_type]
        other = contract.SHEET_NAMES.get(identity.get("record_type", ""), "")
        if sheet in wb.sheetnames:
            ws = wb[sheet]
        elif other in wb.sheetnames:
            # Another type's template: read it, so the header check can say which.
            ws = wb[other]
        elif len(wb.sheetnames) == 1:
            ws = wb[wb.sheetnames[0]]
            notes.append(f"No '{sheet}' sheet: the only sheet, '{ws.title}', was read.")
        else:
            raise UploadRefused(f"The workbook has no '{sheet}' sheet. Use the GHUS "
                                f"{contract.TYPE_LABELS[record_type]} template.")
        rows: list[list[Any]] = []
        for r_index, row in enumerate(ws.iter_rows(max_col=MAX_COLUMNS), 1):
            # A formula cell's value is its formula text ("=..."), never a result:
            # nothing is evaluated, and the row check reports it on its own row.
            values = []
            for cell in row:
                value = cell.value
                if getattr(cell, "data_type", None) == "f" and not str(value).startswith("="):
                    value = f"={value}"
                values.append(value)
            rows.append(values)
            if len(rows) > MAX_ROWS + 1 + 200:      # header + rows + tolerated trailing blanks
                break
        return rows, identity, notes
    finally:
        wb.close()


def _csv_rows(data: bytes) -> tuple[list[list[Any]], list[str]]:
    notes: list[str] = []
    if b"\x00" in data:
        raise UploadRefused("The CSV contains binary data. Save it again as CSV (UTF-8).")
    try:
        text_ = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text_ = data.decode("cp1252", errors="replace")
        notes.append("The CSV is not UTF-8; it was read as Windows-1252. Check accented "
                     "letters in the preview.")
    sample = text_[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample.splitlines()[0] if sample else ",",
                                      delimiters=",;\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    if delimiter != ",":
        notes.append(f"The CSV uses '{delimiter}' between columns; it was read that way.")
    reader = csv.reader(io.StringIO(text_, newline=""), delimiter=delimiter, strict=True)
    rows = []
    try:
        for row in reader:
            rows.append(row[:MAX_COLUMNS])
            if len(rows) > MAX_ROWS + 1 + 200:
                break
    except csv.Error as exc:
        raise UploadRefused(f"The CSV is malformed near line {reader.line_num}: {exc}") from exc
    return rows, notes


def parse_upload(file_name: str, data: bytes, record_type: str) -> ParsedFile:
    name = (file_name or "").strip().replace("\\", "/").rsplit("/", 1)[-1][:200] or "upload"
    extension = name.rsplit(".", 1)[-1].casefold() if "." in name else ""
    if extension not in ("xlsx", "csv"):
        raise UploadRefused("Upload a .xlsx or a .csv file made from the GHUS template.")
    if not data:
        raise UploadRefused("The file is empty.")
    if len(data) > MAX_BYTES:
        raise UploadRefused("The file is larger than 5 MB.")
    identity: dict[str, str] = {}
    if extension == "xlsx":
        raw_rows, identity, notes = _xlsx_rows(data, record_type)
    else:
        raw_rows, notes = _csv_rows(data)
    if not raw_rows:
        raise UploadRefused("The file has no header row.")
    headers = [_cell_text(v) for v in raw_rows[0]]
    while headers and not headers[-1]:
        headers.pop()
    rows: list[dict[str, Any]] = []
    blank = 0
    for offset, raw in enumerate(raw_rows[1:], 2):
        values = [_cell_text(v) for v in raw[:len(headers)]]
        if not any(values):
            blank += 1
            continue
        if len(rows) >= MAX_ROWS:
            raise UploadRefused(f"The file has more than {MAX_ROWS} records. Split it into "
                                "several files.")
        cells: dict[str, str] = {}
        for header, value in zip(headers, values):
            if len(value) > MAX_CELL:
                raise UploadRefused(f"Row {offset}: a value in '{header}' is longer than "
                                    f"{MAX_CELL} characters.")
            cells[header] = value
        rows.append({"row_number": offset, "values": cells})
    return ParsedFile(file_name=name, extension=extension, size=len(data),
                      sha256=hashlib.sha256(data).hexdigest(), headers=headers, rows=rows,
                      template=identity, notes=notes, blank_rows=blank)


def check_headers(parsed: ParsedFile, record_type: str) -> list[str]:
    """Problems with the header row, each a sentence. Empty means it is the template."""
    problems: list[str] = []
    if parsed.template:
        if parsed.template.get("record_type") != record_type:
            other = contract.TYPE_LABELS.get(parsed.template.get("record_type"), "another type")
            problems.append(f"This is the {other} template. Choose {other} above, or "
                            f"use the {contract.TYPE_LABELS[record_type]} template.")
        if parsed.template.get("template_version") != contract.TEMPLATE_VERSION:
            problems.append(f"This file was made from template version "
                            f"{parsed.template.get('template_version')}; the current version "
                            f"is {contract.TEMPLATE_VERSION}. Download the current template.")
    known = contract.by_header(record_type)
    seen: dict[str, str] = {}
    for header in parsed.headers:
        key = contract.header_key(header)
        if not key:
            problems.append("A column has no header. Remove it or restore its header.")
            continue
        if key not in known:
            problems.append(f"Unknown column '{header}'. Headers must not be renamed.")
        if key in seen:
            problems.append(f"The column '{header}' appears twice.")
        seen[key] = header
    missing = [f.template_header for f in contract.fields(record_type)
               if contract.header_key(f.header) not in seen]
    if missing:
        problems.append("Missing columns: " + ", ".join(missing) + ".")
    return problems
