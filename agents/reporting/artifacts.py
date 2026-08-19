"""Template-based artifact builders for the reporting agent."""

from __future__ import annotations

import html
import json
import re
import uuid
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook


@dataclass(frozen=True)
class ArtifactResult:
    artifact_id: str
    artifact_type: str
    path: str
    entity_ids: list[str]


class MissingCitations(ValueError):
    """Raised when a generated artifact has no entity ids behind its figures."""


class ArtifactBuilder:
    def __init__(self, *, template_dir: Path, output_dir: Path):
        self.template_dir = template_dir
        self.output_dir = output_dir

    def build_docx(self, spec: dict[str, Any]) -> ArtifactResult:
        title, body, entity_ids = self._extract_common(spec)
        document_xml = self._render_template(
            "docx_document.xml.template",
            {
                "title": _xml(title),
                "body_paragraphs": "".join(f"<w:p><w:r><w:t>{_xml(p)}</w:t></w:r></w:p>" for p in body),
                "citations": _xml(_citation_text(entity_ids)),
            },
        )
        artifact_id = str(uuid.uuid4())
        path = self._output_path(spec, artifact_id, ".docx")
        self._write_docx(path, document_xml)
        return ArtifactResult(artifact_id, "docx", str(path), entity_ids)

    def build_pptx(self, spec: dict[str, Any]) -> ArtifactResult:
        title, body, entity_ids = self._extract_common(spec)
        slide_xml = self._render_template(
            "pptx_slide.xml.template",
            {
                "title": _xml(title),
                "body": _xml("\n".join(body[:8])),
                "citations": _xml(_citation_text(entity_ids)),
            },
        )
        artifact_id = str(uuid.uuid4())
        path = self._output_path(spec, artifact_id, ".pptx")
        self._write_pptx(path, slide_xml)
        return ArtifactResult(artifact_id, "pptx", str(path), entity_ids)

    def build_xlsx(self, spec: dict[str, Any]) -> ArtifactResult:
        title, body, entity_ids = self._extract_common(spec)
        artifact_id = str(uuid.uuid4())
        path = self._output_path(spec, artifact_id, ".xlsx")
        self._render_template("xlsx_workbook.json.template", {"title": title})

        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"
        summary.append(["Title", title])
        for idx, paragraph in enumerate(body, start=1):
            summary.append([f"Point {idx}", paragraph])
        data = workbook.create_sheet("Data")
        for row in spec.get("rows", []):
            if isinstance(row, dict):
                if data.max_row == 1 and data.max_column == 1 and data["A1"].value is None:
                    data.append(list(row))
                data.append([row.get(key) for key in row])
            elif isinstance(row, list):
                data.append(row)
        citations = workbook.create_sheet("Citations")
        citations.append(["entity_id"])
        for entity_id in entity_ids:
            citations.append([entity_id])
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        return ArtifactResult(artifact_id, "xlsx", str(path), entity_ids)

    def build_pdf(self, spec: dict[str, Any]) -> ArtifactResult:
        title, body, entity_ids = self._extract_common(spec)
        text_body = self._render_template(
            "pdf_body.txt.template",
            {
                "title": title,
                "body": "\n".join(body),
                "citations": _citation_text(entity_ids),
            },
        )
        artifact_id = str(uuid.uuid4())
        path = self._output_path(spec, artifact_id, ".pdf")
        self._write_simple_pdf(path, text_body)
        return ArtifactResult(artifact_id, "pdf", str(path), entity_ids)

    def _extract_common(self, spec: dict[str, Any]) -> tuple[str, list[str], list[str]]:
        title = str(spec.get("title") or "GHUS Report").strip()
        body = _body_lines(spec)
        _validate_item_citations(spec, "figures")
        _validate_item_citations(spec, "tables")
        entity_ids = _entity_ids(spec)
        if not entity_ids:
            raise MissingCitations("generated artifacts must cite at least one entity id")
        return title, body, entity_ids

    def _render_template(self, name: str, values: dict[str, str]) -> str:
        template = (self.template_dir / name).read_text(encoding="utf-8")
        rendered = template
        for key, value in values.items():
            rendered = rendered.replace("{{" + key + "}}", str(value))
        return rendered

    def _output_path(self, spec: dict[str, Any], artifact_id: str, suffix: str) -> Path:
        requested = spec.get("output_path")
        if requested:
            path = Path(str(requested))
            return path if path.is_absolute() else self.output_dir / path
        return self.output_dir / f"{artifact_id}{suffix}"

    @staticmethod
    def _write_docx(path: Path, document_xml: str) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(
                "[Content_Types].xml",
                (
                    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                    '<Default Extension="xml" ContentType="application/xml"/>'
                    '<Override PartName="/word/document.xml" '
                    'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                    "</Types>"
                ),
            )
            archive.writestr(
                "_rels/.rels",
                (
                    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                    '<Relationship Id="rId1" '
                    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
                    'Target="word/document.xml"/></Relationships>'
                ),
            )
            archive.writestr("word/document.xml", document_xml)

    @staticmethod
    def _write_pptx(path: Path, slide_xml: str) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(
                "[Content_Types].xml",
                (
                    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                    '<Default Extension="xml" ContentType="application/xml"/>'
                    '<Override PartName="/ppt/presentation.xml" '
                    'ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>'
                    '<Override PartName="/ppt/slides/slide1.xml" '
                    'ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
                    "</Types>"
                ),
            )
            archive.writestr(
                "_rels/.rels",
                (
                    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                    '<Relationship Id="rId1" '
                    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
                    'Target="ppt/presentation.xml"/></Relationships>'
                ),
            )
            archive.writestr(
                "ppt/presentation.xml",
                (
                    '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
                    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                    "<p:sldIdLst><p:sldId id=\"256\" r:id=\"rId1\"/></p:sldIdLst></p:presentation>"
                ),
            )
            archive.writestr(
                "ppt/_rels/presentation.xml.rels",
                (
                    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                    '<Relationship Id="rId1" '
                    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" '
                    'Target="slides/slide1.xml"/></Relationships>'
                ),
            )
            archive.writestr("ppt/slides/slide1.xml", slide_xml)

    @staticmethod
    def _write_simple_pdf(path: Path, text_body: str) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        lines = _wrap_pdf_lines(text_body)
        content = ["BT", "/F1 10 Tf", "50 760 Td"]
        for idx, line in enumerate(lines[:65]):
            if idx:
                content.append("0 -13 Td")
            content.append(f"({_pdf_text(line)}) Tj")
        content.append("ET")
        stream = "\n".join(content).encode("latin-1", errors="replace")
        objects = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
        ]
        offsets: list[int] = []
        data = bytearray(b"%PDF-1.4\n")
        for idx, obj in enumerate(objects, start=1):
            offsets.append(len(data))
            data.extend(f"{idx} 0 obj\n".encode("ascii"))
            data.extend(obj)
            data.extend(b"\nendobj\n")
        xref = len(data)
        data.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii"))
        for offset in offsets:
            data.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
        data.extend(
            (
                f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\n"
                f"startxref\n{xref}\n%%EOF\n"
            ).encode("ascii")
        )
        path.write_bytes(bytes(data))


def _body_lines(spec: dict[str, Any]) -> list[str]:
    body = spec.get("body") or spec.get("summary") or []
    if isinstance(body, str):
        lines = [line.strip() for line in body.splitlines() if line.strip()]
    elif isinstance(body, list):
        lines = [str(item).strip() for item in body if str(item).strip()]
    else:
        lines = []
    if spec.get("figures"):
        lines.extend(f"Figure: {json.dumps(item, default=str)}" for item in spec["figures"])
    if spec.get("tables"):
        lines.extend(f"Table: {json.dumps(item, default=str)}" for item in spec["tables"])
    return lines or ["No narrative body supplied."]


def _entity_ids(value: Any) -> list[str]:
    seen: set[str] = set()
    ordered: list[str] = []

    def visit(item: Any) -> None:
        if isinstance(item, dict):
            for key, child in item.items():
                if key in {"entity_id", "entity_ids", "citations", "source_entity_ids"}:
                    visit(child)
                elif isinstance(child, (dict, list)):
                    visit(child)
        elif isinstance(item, list):
            for child in item:
                visit(child)
        elif item is not None:
            text = str(item).strip()
            if _looks_like_uuid(text) and text not in seen:
                seen.add(text)
                ordered.append(text)

    visit(value)
    return ordered


def _validate_item_citations(spec: dict[str, Any], key: str) -> None:
    items = spec.get(key) or []
    if not isinstance(items, list):
        raise MissingCitations(f"{key} must be a list when supplied")
    for idx, item in enumerate(items):
        if not _entity_ids(item):
            raise MissingCitations(f"{key}[{idx}] must cite at least one entity id")


def _looks_like_uuid(value: str) -> bool:
    return bool(re.fullmatch(r"[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}", value))


def _citation_text(entity_ids: Iterable[str]) -> str:
    return "Cited entity ids: " + ", ".join(entity_ids)


def _xml(value: str) -> str:
    return html.escape(str(value), quote=False)


def _pdf_text(value: str) -> str:
    return str(value).replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _wrap_pdf_lines(value: str, width: int = 92) -> list[str]:
    lines: list[str] = []
    for raw in value.splitlines():
        text = raw.strip()
        while len(text) > width:
            lines.append(text[:width])
            text = text[width:]
        lines.append(text)
    return lines
