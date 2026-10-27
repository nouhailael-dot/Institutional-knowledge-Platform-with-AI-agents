"""Data Feed: templates, safe parsing, row checks, preview, confirmation, Review.

Most tests need nothing but this checkout. The end-to-end tests stage into a
disposable local database (DATAFEED_TEST_DATABASE_URL ending in
``_datafeed_test``) through the real backend command, and remove what they
staged afterwards.
"""

from __future__ import annotations

import io
import json
import os
import time
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from src.datafeed import contract, parse, reference, templates, validate
from src.datafeed.service import DataFeedError, DataFeedService

VOCAB = reference.vocabularies(None)


def filled_xlsx(record_type: str, rows: list[dict[str, str]]) -> bytes:
    """The official template with rows typed into it, by header (no required mark)."""
    wb = load_workbook(io.BytesIO(templates.build_xlsx(record_type, VOCAB)))
    ws = wb[contract.SHEET_NAMES[record_type]]
    columns = {contract.header_key(c.value): c.column for c in ws[1]}
    for r, row in enumerate(rows, 2):
        for header, value in row.items():
            ws.cell(row=r, column=columns[contract.header_key(header)], value=value)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def csv_bytes(record_type: str, rows: list[dict[str, str]], *, sep=",", encoding="utf-8-sig"):
    headers = [f.template_header for f in contract.fields(record_type)]
    by_key = {contract.header_key(h): i for i, h in enumerate(headers)}
    lines = [sep.join(headers)]
    for row in rows:
        cells = [""] * len(headers)
        for header, value in row.items():
            cells[by_key[contract.header_key(header)]] = value
        lines.append(sep.join(f'"{c}"' if sep in c or '"' in c else c for c in cells))
    return ("\r\n".join(lines) + "\r\n").encode(encoding)


ACTOR = {"Official name": "Prairie Water Lab Inc.", "Website": "https://prairiewaterlab-test.org",
         "Actor Category": "Company / Corporate Actor", "Category Type": "SME",
         "Description": "A company that builds water sensors for farms.", "Sector": "Water",
         "City": "Wichita", "State": "Kansas",
         "Source URL": "https://prairiewaterlab-test.org/about"}


# ------------------------------------------------------------------ templates
def test_templates_follow_the_contract_and_carry_their_identity():
    for record_type in contract.RECORD_TYPES:
        wb = load_workbook(io.BytesIO(templates.build_xlsx(record_type, VOCAB)))
        assert wb.sheetnames == ["Instructions", contract.SHEET_NAMES[record_type], "Field Guide",
                                 "Examples", "Lists"]
        assert wb["Lists"].sheet_state == "veryHidden"
        assert [c.value for c in wb["Lists"][1]][:3] == ["ghus-datafeed", record_type,
                                                         contract.TEMPLATE_VERSION]
        data = wb[contract.SHEET_NAMES[record_type]]
        headers = [c.value for c in data[1]]
        assert headers == [f.template_header for f in contract.fields(record_type)]
        assert not any(any(c.value for c in row) for row in data.iter_rows(min_row=2))
        assert all(r[0].value.startswith("Column") or r[0].value for r in wb["Field Guide"].iter_rows())
        names = [r[0].value for r in wb["Examples"].iter_rows(min_row=3)]
        assert names and all(n.startswith("EXAMPLE - ") for n in names)
        csv_headers = templates.build_csv(record_type).decode("utf-8-sig").strip().split(",")
        assert csv_headers == headers


def test_required_columns_are_marked_and_internal_columns_never_offered():
    actor = [f.template_header for f in contract.fields("actor")]
    assert "Official name *" in actor and "Actor Category *" in actor and "Sector *" in actor
    every = " ".join(f.key for t in contract.RECORD_TYPES for f in contract.fields(t))
    for internal in ("actor_id", "created_at", "updated_at", "verification", "run_id",
                     "status", "actor_type", "pointer", "rankings", "published_stage_or_trl"):
        assert internal not in every.split()


def test_dropdowns_cover_every_controlled_column_and_type_follows_category():
    wb = load_workbook(io.BytesIO(templates.build_xlsx("actor", VOCAB)))
    rules = wb["Actors"].data_validations.dataValidation
    formulas = [r.formula1 for r in rules]
    assert any(f.startswith("INDIRECT(") for f in formulas)          # Category Type
    assert not any(f.startswith("=") for f in formulas)
    names = set(wb.defined_names)
    for category in VOCAB["actor_categories"]:
        assert templates.range_name(category["label"]) in names
    offered = [f for f in contract.fields("actor") if f.kind == "enum"
               and (f.vocab == "category_types" or VOCAB.get(f.vocab))]
    assert len(rules) == len(offered)          # topics come from the database only


# ------------------------------------------------------------------ parsing
def test_a_filled_template_and_its_csv_read_the_same():
    from_xlsx = parse.parse_upload("a.xlsx", filled_xlsx("actor", [ACTOR]), "actor")
    from_csv = parse.parse_upload("a.csv", csv_bytes("actor", [ACTOR]), "actor")
    assert from_xlsx.template["template_version"] == contract.TEMPLATE_VERSION
    assert parse.check_headers(from_xlsx, "actor") == []
    assert parse.check_headers(from_csv, "actor") == []
    assert from_xlsx.rows[0]["row_number"] == 2
    picked = {k: v for k, v in from_xlsx.rows[0]["values"].items() if v}
    assert picked == {k: v for k, v in from_csv.rows[0]["values"].items() if v}


def test_semicolon_and_windows_csv_are_read_and_said():
    parsed = parse.parse_upload("a.csv", csv_bytes("actor", [{**ACTOR, "City": "Mérida"}],
                                                   sep=";", encoding="cp1252"), "actor")
    assert any("';'" in n for n in parsed.notes)
    assert any("Windows-1252" in n for n in parsed.notes)
    assert parsed.rows[0]["values"]["City"] == "Mérida"


@pytest.mark.parametrize("name, data, message", [
    ("a.txt", b"x", ".xlsx or a .csv"),
    ("a.csv", b"", "empty"),
    ("a.csv", b"x" * (parse.MAX_BYTES + 1), "5 MB"),
    ("a.xlsx", b"not a zip", "not a valid .xlsx"),
    ("a.csv", b"a,b\x00,c", "binary"),
], ids=["extension", "empty", "too-large", "not-a-workbook", "binary-csv"])
def test_unsafe_or_wrong_files_are_refused(name, data, message):
    with pytest.raises(parse.UploadRefused, match=message):
        parse.parse_upload(name, data, "actor")


def _zip(parts: dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, body in parts.items():
            archive.writestr(name, body)
    return out.getvalue()


def test_macros_external_links_and_compression_bombs_are_refused():
    base = {"[Content_Types].xml": b"<x/>", "xl/workbook.xml": b"<x/>"}
    with pytest.raises(parse.UploadRefused, match="macros"):
        parse.check_container(_zip({**base, "xl/vbaProject.bin": b"x"}))
    with pytest.raises(parse.UploadRefused, match="macros"):
        parse.check_container(_zip({**base, "xl/externalLinks/externalLink1.xml": b"x"}))
    with pytest.raises(parse.UploadRefused, match="compressed suspiciously|60 MB"):
        parse.check_container(_zip({**base, "xl/worksheets/sheet1.xml": b"A" * (8 * 1024 * 1024)}))


def test_formulas_are_never_evaluated_and_are_reported_on_their_row():
    data = filled_xlsx("actor", [{**ACTOR, "Official name": '=HYPERLINK("http://x","y")'}])
    parsed = parse.parse_upload("a.xlsx", data, "actor")
    row = parsed.rows[0]
    result = validate.RowChecker("actor", VOCAB).check(row["row_number"], row["values"])
    assert any("starts like a formula" in e for e in result["errors"])
    assert not any("is required" in e for e in result["errors"])
    assert parse.formula_like("@SUM(A1)") and parse.formula_like("+cmd")
    assert not parse.formula_like("-12.5") and not parse.formula_like("Plain text")


def test_headers_must_match_the_selected_template_version():
    parsed = parse.parse_upload("e.xlsx", filled_xlsx("event", []), "actor")
    problems = parse.check_headers(parsed, "actor")
    assert any("Events template" in p for p in problems)
    renamed = csv_bytes("actor", [ACTOR]).replace(b"Official name *", b"Org name")
    problems = parse.check_headers(parse.parse_upload("a.csv", renamed, "actor"), "actor")
    assert "Unknown column 'Org name'. Headers must not be renamed." in problems
    assert any(p.startswith("Missing columns: Official name *") for p in problems)
    old = parse.parse_upload("a.xlsx", filled_xlsx("actor", [ACTOR]), "actor")
    old.template["template_version"] = "2026-01-01.1"
    assert any("Download the current template" in p for p in parse.check_headers(old, "actor"))


# ------------------------------------------------------------------ row checks
def check(record_type, row, number=2):
    values = {h: "" for h in (f.template_header for f in contract.fields(record_type))}
    for header, value in row.items():
        values[next(h for h in values if contract.header_key(h) == contract.header_key(header))] = value
    return validate.RowChecker(record_type, VOCAB).check(number, values)


def test_a_valid_actor_becomes_the_staged_payload():
    result = check("actor", {**ACTOR, "Recent activity": "Pilot with K-State",
                             "Recent activity date": "2026-05-14",
                             "Company: ownership and size": "Private, 25 employees"})
    assert result["errors"] == []
    p = result["payload"]
    assert (p["actor_category"], p["category_type"], p["state_code"], p["primary_sector"]) == \
        ("company", "small_mid_size", "KS", "WATER")
    assert p["country"] == "United States"                     # implied by a US state, and said
    assert any("Country was empty" in w for w in result["warnings"])
    assert p["light_pull_values"] == {
        "recent_activity": {"description": "Pilot with K-State", "date": "2026-05-14"},
        "ownership_and_size": "Private, 25 employees"}


def test_category_and_type_must_belong_together():
    result = check("actor", {**ACTOR, "Actor Category": "University / Academic Institution"}, 12)
    assert ("Row 12: Actor Category is University / Academic Institution, but Category Type "
            "'SME' belongs to Company / Corporate Actor.") in result["errors"]
    result = check("actor", {**ACTOR, "Category Type": "Not Applicable"})
    assert any("needs a one-line Category note" in e for e in result["errors"])
    result = check("actor", {**ACTOR, "Category Type": ""})
    assert result["errors"] == [] and "category_type" not in result["payload"]
    assert any("Pending Classification" in w for w in result["warnings"])


def test_required_fields_types_and_lists():
    result = check("actor", {"Website": "www.no-scheme.org", "Actor Category": "Bakery",
                             "Sector": "Space"})
    errors = " ".join(result["errors"])
    assert "Official name is required" in errors
    assert "Website is not a web address" in errors
    assert "Actor Category 'Bakery' is not an allowed value" in errors
    assert "Sector 'Space' is not an allowed value" in errors


def test_geography_is_normalised_and_the_region_is_never_trusted():
    assert check("actor", {**ACTOR, "State": "ks"})["payload"]["state_code"] == "KS"
    result = check("actor", {**ACTOR, "State": "TX", "Region": "Northeast"})
    assert any("does not match the State TX" in e and "South Central / Texas Gulf" in e
               for e in result["errors"])
    assert check("actor", {**ACTOR, "State": "TX", "Region": "South Central / Texas Gulf"})[
        "errors"] == []
    result = check("actor", {**ACTOR, "State": "Atlantis"})
    assert any("is not a US state" in e for e in result["errors"])
    result = check("actor", {**ACTOR, "State": "", "Country": "Morocco"})
    assert result["errors"] == [] and any("outside the US" in w for w in result["warnings"])


def test_light_pull_boxes_belong_to_their_category_and_activity_needs_a_date():
    result = check("actor", {**ACTOR, "University: public or private": "Public"})
    assert any("is only for University / Academic Institution" in e for e in result["errors"])
    result = check("actor", {**ACTOR, "Recent activity": "Opened a lab"})
    assert any("needs its date" in e for e in result["errors"])
    long = " ".join(["One sentence."] * 4)
    assert any("1-3 sentences" in e for e in check("actor", {**ACTOR, "Description": long})["errors"])


def test_classification_evidence_columns_are_required_where_the_rules_need_them():
    assert any("Description is required" in e
               for e in check("actor", {**ACTOR, "Description": ""})["errors"])
    startup = {**ACTOR, "Actor Category": "Startup / Emerging Technology Company",
               "Category Type": "University Spinout"}
    assert any("published development stage" in e for e in check("actor", startup)["errors"])
    ok = check("actor", {**startup, "Startup: development stage": "Pilot with two utilities"})
    assert ok["errors"] == [] and ok["payload"]["development_stage"] == "Pilot with two utilities"
    assert any("only for Startup" in e for e in check(
        "actor", {**ACTOR, "Startup: development stage": "Seed"})["errors"])
    lab = {**ACTOR, "Actor Category": "National Laboratory", "Category Type": "Government-Direct"}
    assert any("formal designation" in e for e in check("actor", lab)["errors"])


def test_parent_needs_its_relationship_and_example_rows_are_refused():
    result = check("actor", {**ACTOR, "Parent organization": "University of Arizona"})
    assert any("Relationship to parent is required" in e for e in result["errors"])
    result = check("actor", {**ACTOR, "Official name": "EXAMPLE - Prairie Water Lab"})
    assert any("example row" in e for e in result["errors"])


def test_event_dates_and_identity():
    event = {"Event name": "Great Plains Water Summit", "Event type": "Summit",
             "Start date": "2027-04-02", "End date": "2027-04-03",
             "Official event URL": "https://gpws-test.org", "Format": "Online"}
    result = check("event", event)
    assert result["errors"] == []
    assert result["payload"]["location"] == "Online"
    assert any("no confirmed host" in w for w in result["warnings"])
    result = check("event", {**event, "End date": "2027-04-01"})
    assert any("is before the Start date" in e for e in result["errors"])
    result = check("event", {**event, "Start date": "04/02/2027"})
    assert any("is not a date" in e for e in result["errors"])
    result = check("event", {**event, "Event type": "Party"})
    assert any("Event type 'Party' is not an allowed value" in e for e in result["errors"])


def test_people_need_an_organisation_a_checkable_link_and_a_real_linkedin():
    person = {"Full name": "Ana Ruiz", "Organization": "University of Arizona",
              "Public LinkedIn URL": "https://www.linkedin.com/in/ana-test",
              "Expertise": "groundwater; aquifer recharge"}
    result = check("person", person)
    assert result["errors"] == []
    assert result["payload"]["research_area"] == "groundwater; aquifer recharge"
    result = check("person", {**person, "Public LinkedIn URL": "https://example.org/ana"})
    assert any("not a linkedin.com address" in e for e in result["errors"])
    result = check("person", {"Full name": "Ana Ruiz", "Organization": "University of Arizona"})
    assert any("so the record can be checked" in e for e in result["errors"])
    assert any("Profile incomplete" in w for w in check("person", {
        **person, "Public LinkedIn URL": "", "Source URL": "https://a.edu/x"})["warnings"])


def test_duplicate_rows_in_one_file_are_named():
    rows = [check("actor", ACTOR, 2), check("actor", {**ACTOR, "Website":
                                                      "https://www.prairiewaterlab-test.org/x"}, 7)]
    validate.file_duplicates("actor", rows)
    assert rows[1]["errors"] == ["Row 7: same record as row 2 (website)."]


# ------------------------------------------------------------------ service
class FakeBridge:
    configured = True
    database_url = ""

    def __init__(self):
        self.checked, self.staged = [], []

    def check(self, batch):
        self.checked.append(batch)
        return {"status": "checked", "rows": [
            {"row_number": r["row_number"], "errors": [], "warnings": [],
             "outcome": "would_stage"} for r in batch["rows"]], "already_imported_run_id": None}

    def stage(self, batch):
        self.staged.append(batch)
        return {"status": "staged", "run_id": "00000000-0000-0000-0000-00000000abcd",
                "staged": len(batch["rows"]), "already_awaiting_review": 0,
                "rows": [{"row_number": r["row_number"], "outcome": "staged"}
                         for r in batch["rows"]]}


@pytest.fixture
def service(tmp_path):
    return DataFeedService(bridge=FakeBridge(), connect=None, preview_dir=tmp_path / "previews")


def test_preview_then_explicit_confirmation_stages_only_valid_rows(service):
    bad = {**ACTOR, "Official name": "", "Website": "https://other-test.org"}
    preview = service.validate_upload("actor", "kansas.xlsx", filled_xlsx("actor", [ACTOR, bad]))
    assert preview["counts"]["total"] == 2 and preview["counts"]["needs_correction"] == 1
    assert preview["can_confirm"] and preview["file"]["name"] == "kansas.xlsx"
    assert len(preview["file"]["sha256"]) == 64
    assert service.bridge.staged == []                               # nothing yet
    with pytest.raises(DataFeedError, match="would be left out"):
        service.confirm(preview["preview_id"], uploaded_by="Test Uploader", attest_source=True, import_name="Kansas", accept_left_out=False)
    result = service.confirm(preview["preview_id"], uploaded_by="Test Uploader", attest_source=True, import_name="Kansas", accept_left_out=True)
    assert result["staged"] == 1 and result["left_out"] == 1
    batch = service.bridge.staged[0]
    assert [r["row_number"] for r in batch["rows"]] == [2]
    assert batch["rows_total"] == 2 and batch["rows_invalid"] == 1
    assert batch["uploaded_by"] == "Test Uploader" and batch["file_name"] == "kansas.xlsx"
    assert batch["admission"] == "review" and batch["attestation"]["source"] is True
    assert "Compass import" in batch["attestation"]["text"]
    assert batch["rows"][0]["original"]["Official name *"] == "Prairie Water Lab Inc."
    with pytest.raises(DataFeedError, match="expired"):
        service.preview(preview["preview_id"])                         # deleted once staged


def test_the_uploader_is_named_confirms_the_source_and_attests_for_verified(service):
    preview = service.validate_upload("actor", "a.csv", csv_bytes("actor", [ACTOR]))
    pid = preview["preview_id"]
    with pytest.raises(DataFeedError, match="Give your name"):
        service.confirm(pid, import_name="K", accept_left_out=False, attest_source=True)
    with pytest.raises(DataFeedError, match="GHUS source"):
        service.confirm(pid, import_name="K", accept_left_out=False, uploaded_by="Ana")
    with pytest.raises(DataFeedError, match="authorized GHUS staff"):
        service.confirm(pid, import_name="K", accept_left_out=False, uploaded_by="Ana",
                        attest_source=True, admission="verified")
    done = service.confirm(pid, import_name="K", accept_left_out=False, uploaded_by="Ana",
                           attest_source=True, admission="verified", attest_authorized=True)
    batch = service.bridge.staged[-1]
    assert batch["admission"] == "verified" and batch["attestation"]["authorized"] is True
    assert done["admission"] == "verified" and done["uploaded_by"] == "Ana"
    events = service.validate_manual("event", [{"name": "Water Summit", "event_type": "Summit",
                                                "start_date": "2027-04-02",
                                                "website": "https://ws-test.org"}])
    with pytest.raises(DataFeedError, match="Only organizations"):
        service.confirm(events["preview_id"], import_name="E", accept_left_out=False,
                        uploaded_by="Ana", attest_source=True, admission="verified",
                        attest_authorized=True)


def test_a_preview_needs_a_name_and_expires(service, monkeypatch):
    preview = service.validate_upload("actor", "a.csv", csv_bytes("actor", [ACTOR]))
    with pytest.raises(DataFeedError, match="name"):
        service.confirm(preview["preview_id"], uploaded_by="Test Uploader", attest_source=True, import_name=" ", accept_left_out=False)
    path = service.preview_dir / f"{preview['preview_id']}.json"
    old = time.time() - 25 * 3600
    os.utime(path, (old, old))
    with pytest.raises(DataFeedError, match="expired"):
        service.preview(preview["preview_id"])
    with pytest.raises(DataFeedError, match="Unknown preview"):
        service.preview("../../etc/passwd")


def test_manual_entry_uses_the_same_checks(service):
    preview = service.validate_manual("event", [{
        "name": "Great Plains Water Summit", "event_type": "Summit", "start_date": "2027-04-02",
        "website": "https://gpws-test.org"}])
    assert preview["method"] == "manual" and preview["counts"]["valid"] == 1
    with pytest.raises(DataFeedError, match="Unknown fields"):
        service.validate_manual("event", [{"event_id": "x"}])
    with pytest.raises(DataFeedError, match="At most"):
        service.validate_manual("event", [{}] * 51)


def test_wrong_columns_stop_at_the_header(service):
    preview = service.validate_upload("person", "a.xlsx", filled_xlsx("actor", [ACTOR]))
    assert preview["header_problems"] and not preview["rows"] and not preview["can_confirm"]


def test_unconfigured_staging_still_checks_but_cannot_send(tmp_path):
    bridge = FakeBridge()
    bridge.configured = False
    service = DataFeedService(bridge=bridge, connect=None, preview_dir=tmp_path)
    preview = service.validate_upload("actor", "a.csv", csv_bytes("actor", [ACTOR]))
    assert preview["counts"]["valid"] == 1 and not preview["can_confirm"]
    assert any("staging is not configured" in n for n in preview["notes"])


def test_routes(monkeypatch, tmp_path):
    from fastapi.testclient import TestClient

    from backend import datafeed_routes
    from backend.app import app

    service = DataFeedService(bridge=FakeBridge(), connect=None, preview_dir=tmp_path)
    monkeypatch.setattr(datafeed_routes, "datafeed_service", lambda: service)
    client = TestClient(app)
    catalog = client.get("/api/datafeed/catalog").json()
    assert catalog["template_version"] == contract.TEMPLATE_VERSION
    xlsx = client.get("/api/datafeed/templates/actor.xlsx")
    assert xlsx.status_code == 200 and "GHUS_Actors_Template.xlsx" in xlsx.headers[
        "content-disposition"]
    assert client.get("/api/datafeed/templates/event.csv").content.decode("utf-8-sig").startswith(
        "Event name *")
    assert client.get("/api/datafeed/templates/hub.xlsx").status_code == 400
    response = client.post("/api/datafeed/validate", data={"record_type": "actor"},
                           files={"file": ("a.csv", csv_bytes("actor", [ACTOR]), "text/csv")})
    preview = response.json()
    assert response.status_code == 200 and preview["counts"]["valid"] == 1
    done = client.post(f"/api/datafeed/previews/{preview['preview_id']}/confirm",
                       json={"import_name": "Route test", "uploaded_by": "Test Uploader",
                             "attest_source": True}).json()
    assert done["staged"] == 1
    assert client.post("/api/datafeed/validate", data={"record_type": "actor"},
                       files={"file": ("a.txt", b"x", "text/plain")}).status_code == 400


# ------------------------------------------------------------------ end to end (isolated DB)
URL = os.environ.get("DATAFEED_TEST_DATABASE_URL", "")
BACKEND = Path(os.environ.get("GHUS_DATAFEED_PROJECT_ROOT")
               or Path(__file__).resolve().parents[1] / "ghus-backend")
needs_db = pytest.mark.skipif(
    not URL or "127.0.0.1:55432" not in URL or not URL.endswith("_datafeed_test")
    or not (BACKEND / "agents" / "datafeed" / "stage_import.py").is_file(),
    reason="DATAFEED_TEST_DATABASE_URL (*_datafeed_test) and the backend checkout are required")


@pytest.fixture
def live(tmp_path):
    import psycopg

    from src.datafeed.bridge import DataFeedBridge

    bridge = DataFeedBridge(project_root=BACKEND, database_url=URL)
    service = DataFeedService(bridge=bridge, preview_dir=tmp_path)
    reference.clear_cache()
    runs: list[str] = []
    yield service, runs
    conninfo = URL.replace("postgresql+psycopg://", "postgresql://")
    with psycopg.connect(conninfo, autocommit=True) as c:      # remove what the test staged
        for run in runs:
            sources = [r[0] for r in c.execute(
                "SELECT q.source_id FROM person_review_queue q JOIN staged_record s "
                "USING (staged_record_id) WHERE s.run_id = %s::uuid", (run,)).fetchall()]
            # The run's staged rows and queue rows go with it (ON DELETE CASCADE).
            c.execute("DELETE FROM scrape_run WHERE run_id = %s::uuid", (run,))
            c.execute("DELETE FROM source WHERE source_id = ANY(%s)", (sources,))


@needs_db
def test_end_to_end_upload_matches_stages_and_appears_in_review(live):
    service, runs = live
    suffix = str(int(time.time() * 1000))
    rows = [{**ACTOR, "Official name": f"Prairie Water Lab {suffix} Inc.",
             "Website": f"https://pwl-{suffix}.org", "Source URL": f"https://pwl-{suffix}.org/a"},
            {"Official name": "ADM", "Website": "https://www.adm.com",
             "Actor Category": "Company / Corporate Actor", "Category Type": "Public Company",
             "Description": "An agricultural processing and food ingredient company.",
             "Sector": "Agriculture", "State": "IL"},
            {**ACTOR, "Official name": f"Water Unit {suffix}", "Website": f"https://wu-{suffix}.org",
             "Actor Category": "Research Center / Institute / Laboratory", "Category Type": "",
             "Parent organization": "3M Research Laboratories", "Relationship to parent": "Unit of"}]
    data = filled_xlsx("actor", rows)
    preview = service.validate_upload("actor", f"e2e-{suffix}.xlsx", data)
    statuses = [r["status"] for r in preview["rows"]]
    # ADM may already wait in Review from an earlier import: then it is not sent again.
    assert statuses[0] == "new" and statuses[2] == "new", preview["rows"]
    assert statuses[1] in ("existing", "waiting")
    assert preview["backend_checked"]
    assert preview["rows"][1]["match"]["name"] == "ADM (Archer Daniels Midland)"
    assert preview["rows"][2]["payload"]["parent_connection"]["relationship_type"] == "unit_of"
    result = service.confirm(preview["preview_id"], uploaded_by="Test Uploader", attest_source=True, import_name=f"E2E {suffix}",
                             accept_left_out=False)
    runs.append(result["run_id"])
    assert result["status"] == "staged" and result["staged"] == (3 if statuses[1] == "existing" else 2)

    # The same file again is not staged twice.
    again = service.validate_upload("actor", f"e2e-{suffix}.xlsx", data)
    assert again["already_imported_run_id"] == result["run_id"] and not again["can_confirm"]
    assert {r["status"] for r in again["rows"]} <= {"waiting", "existing"}

    # History and per-record status.
    item = next(i for i in service.history() if i["run_id"] == result["run_id"])
    assert (item["import_name"], item["counts"]["pending"], item["status"]) == (
        f"E2E {suffix}", result["staged"], "Awaiting Review")
    assert {r["status"] for r in service.import_records(result["run_id"])} == {"Awaiting Review"}

    # Review lists them under their own Data Feed run, with the match offered.
    from src.review.postgres_source import PostgresReviewSource

    source = PostgresReviewSource(URL.replace("+psycopg", ""))
    candidates, _ = source.fetch_candidates()
    mine = [c for c in candidates if c["source_run_id"] == f"datafeed:{result['run_id']}"]
    assert len(mine) == result["staged"]
    assert all(c["source_description"] == f"Data Feed import: E2E {suffix} · e2e-{suffix}.xlsx"
               for c in mine)
    for candidate in mine:
        assert candidate["payload"]["_record_origin"] == "ghus_uploaded_doc"
    adm = next((c for c in mine if c["payload"]["name"] == "ADM"), None)
    if adm is not None:
        assert adm["agent_metadata"]["top_matches"][0]["name"] == "ADM (Archer Daniels Midland)"


@needs_db
def test_end_to_end_people_and_events(live):
    service, runs = live
    suffix = str(int(time.time() * 1000))
    people = service.validate_manual("person", [{
        "name": f"Ana Ruiz {suffix}", "role": "Professor of Hydrology",
        "affiliation": "University of Arizona", "linkedin_url": f"https://www.linkedin.com/in/ana-{suffix}",
        "expertise": "groundwater", "sectors": "Water"}])
    assert people["counts"]["valid"] == 1, people["rows"]
    staged = service.confirm(people["preview_id"], uploaded_by="Test Uploader", attest_source=True, import_name=f"People {suffix}",
                             accept_left_out=False)
    runs.append(staged["run_id"])
    events = service.validate_upload("event", "ev.csv", csv_bytes("event", [{
        "Event name": f"Water Summit {suffix}", "Event type": "Summit", "Start date": "2027-04-02",
        "End date": "2027-04-01", "Official event URL": "https://ws-test.org"}, {
        "Event name": f"Water Forum {suffix}", "Event type": "Forum", "Start date": "2027-05-02",
        "Official event URL": "https://wf-test.org"}]))
    assert [r["status"] for r in events["rows"]] == ["invalid", "new"]
    staged = service.confirm(events["preview_id"], uploaded_by="Test Uploader", attest_source=True, import_name=f"Events {suffix}",
                             accept_left_out=True)
    runs.append(staged["run_id"])
    item = next(i for i in service.history() if i["run_id"] == staged["run_id"])
    assert item["rows_invalid"] == 1 and item["staged"] == 1 and item["rows_total"] == 2


def test_json_preview_is_never_served():
    # Previews live under .datafeed/, which the static mount (frontend/) never serves.
    from src.datafeed import service as module

    assert "frontend" not in str(module.ROOT / ".datafeed")
    assert json.loads(json.dumps(contract.as_payload()))["template_version"]


# ------------------------------------------------------------------ Review labelling
def test_review_names_a_data_feed_record_honestly():
    from src.review.postgres_source import PostgresReviewSource
    from src.review.store import derive_review_reason

    assert derive_review_reason({"data_feed": {"import_name": "K"}, "score_band": "review",
                                 "source_table": "search_candidate"}) == "data_feed"
    assert derive_review_reason({"data_feed": {}, "top_matches": [{"entity_id": "a"}],
                                 "review_kind": "rubric_band"}) == "possible_duplicate"
    assert derive_review_reason({"source_table": "search_candidate",
                                 "score_band": "review"}) == "borderline_score"

    row = {"run_source": "ghus_data_feed", "run_id": "r1", "run_import_name": "Kansas",
           "run_file_name": None}
    payload = {"data_feed": {"method": "manual", "match": {"kind": "existing", "entity_id": "a",
                                                           "name": "ADM", "score": 1.0}}}
    record = {"agent_metadata": {"facts": [{"fact_key": "website", "read_from": "discovery_page",
                                            "source_origin": "newly_found"}]},
              "evidence": [{"url": "https://x.org"}], "source_run_id": "discovery:r1"}
    out = PostgresReviewSource._datafeed(row, payload, record)
    assert out["source_run_id"] == "datafeed:r1"
    assert out["source_description"] == "Data Feed import: Kansas · manual entry"
    assert out["agent_metadata"]["top_matches"][0]["name"] == "ADM"
    assert out["duplicate_state"] == "possible_duplicate"
    assert out["agent_metadata"]["facts"][0]["source_origin"] == "ghus_manual"
    # A discovery candidate is left exactly as it was.
    untouched = {"agent_metadata": {}, "source_run_id": "discovery:r2"}
    assert PostgresReviewSource._datafeed({"run_source": "open_web_search"}, {}, untouched) \
        is untouched


def test_derived_values_are_not_reported_missing():
    result = check("actor", ACTOR)                    # no Country typed; a US state implies it
    assert "Country" not in result["missing"]
