"""Offline V2 actor contract and geography checks."""
import unittest
from src.map_agent.actor_profile import CATEGORIES, CATEGORY_TYPES, REGIONS, normalize_actor_profile, region_for
from src.map_agent.entity_schemas import submit_entities_tool
from src.map_agent.export import _field_value
from src.map_agent.dedup import deduplicate

URL='https://example.edu/project'
class ActorProfileTests(unittest.TestCase):
    def test_category_schema_and_manual_fields(self):
        props=submit_entities_tool('actor')['input_schema']['properties']['entities']['items']['properties']
        self.assertEqual(props['actor_type']['enum'],CATEGORIES)
        self.assertEqual(len(CATEGORIES),12)
        for field in ('record_status','ghus_notes','verification_history','region'):
            self.assertNotIn(field,props)

    def test_retained_fields_and_relationship_sources(self):
        actor={'name':'Example','actor_type':CATEGORIES[0], 'country':'US','state':'VA',
               'sector':'Water','multi_location':True,
               'category_details':[{'label':'Program','value':'Water research'}],
               'actor_relationships':[{'organization_name':'Parent','relationship':'unit of','source_url':URL},
                   {'organization_name':'Invented','relationship':'unit of','source_url':'https://invented.example'}]}
        normalize_actor_profile(actor,{URL:{'kind':'page'}})
        self.assertEqual(actor['sector'],'Water')
        self.assertEqual(actor['region'],'Mid-Atlantic')
        self.assertTrue(actor['multi_location'])
        self.assertEqual(len(actor['category_details']),1)
        self.assertEqual(len(actor['actor_relationships']),1)
        for field in ('field_evidence','record_status','hub_status','um6p_ocp_status'):
            self.assertNotIn(field,actor)

    def test_categories_use_the_intern_guide_names(self):
        for name in ('Innovation & Entrepreneurship (E&I) Platform','Investor / Funding Platform',
                     'Association / Consortium / Network'):
            self.assertIn(name,CATEGORIES)

    def test_a_replayed_old_category_name_is_renamed_not_flagged(self):
        actor={'name':'Example Accelerator',
               'actor_type':'Innovation Platform / Incubator / Accelerator / Venture Studio'}
        normalize_actor_profile(actor,{})
        self.assertEqual(actor['actor_type'],'Innovation & Entrepreneurship (E&I) Platform')
        self.assertNotIn('category_note',actor)

    def test_an_unknown_category_gets_a_note(self):
        actor={'name':'Example','actor_type':'Think tank'}
        normalize_actor_profile(actor,{})
        self.assertEqual(actor['category_note'],'Category needs review: Think tank')

    def test_every_category_has_the_guide_types_plus_not_applicable(self):
        self.assertEqual(list(CATEGORY_TYPES),CATEGORIES)
        for category,types in CATEGORY_TYPES.items():
            self.assertEqual(types[-1],'Not Applicable',category)
            self.assertEqual(len(types),len(set(types)),category)
        self.assertEqual(len(CATEGORY_TYPES['University / Academic Institution']),11)
        self.assertIn('Medical School',CATEGORY_TYPES['University / Academic Institution'])
        self.assertIn('SME',CATEGORY_TYPES['Company / Corporate Actor'])
        self.assertIn('Diaspora Network',CATEGORY_TYPES['Association / Consortium / Network'])

    def test_schema_requires_a_category_type_from_the_lists(self):
        item=submit_entities_tool('actor')['input_schema']['properties']['entities']['items']
        self.assertIn('category_type',item['required'])
        self.assertEqual(set(item['properties']['category_type']['enum']),
                         {t for types in CATEGORY_TYPES.values() for t in types})

    def test_a_type_from_the_category_is_kept(self):
        actor={'name':'Dept. of Plant Sciences','actor_type':'University / Academic Institution',
               'category_type':'Department'}
        normalize_actor_profile(actor,{})
        self.assertEqual(actor['category_type'],'Department')
        self.assertNotIn('category_note',actor)

    def test_a_type_from_another_category_becomes_not_applicable_with_a_note(self):
        actor={'name':'Example Co','actor_type':'Company / Corporate Actor','category_type':'Medical School'}
        normalize_actor_profile(actor,{})
        self.assertEqual(actor['category_type'],'Not Applicable')
        self.assertIn('Medical School',actor['category_note'])

    def test_a_missing_type_is_not_classified(self):
        actor={'name':'Old record','actor_type':'National Laboratory'}
        normalize_actor_profile(actor,{})
        self.assertEqual(actor['category_type'],'Not classified')

    def test_merging_duplicates_keeps_each_type_with_its_category(self):
        lab={'_entity_type':'actor','name':'Example Lab','website':'https://lab.example.edu',
             'actor_type':'University / Academic Institution','category_type':'Research Lab / Center'}
        longer={'_entity_type':'actor','name':'Example Lab','website':'https://lab.example.edu',
                'actor_type':'Research Center / Institute / Laboratory','category_type':'Independent'}
        merged=deduplicate([lab,longer])[0]
        self.assertEqual((merged['actor_type'],merged['category_type']),
                         ('Research Center / Institute / Laboratory','Independent'))
        unclassified={**lab,'category_type':'Not classified'}
        self.assertEqual(deduplicate([unclassified,lab])[0]['category_type'],'Research Lab / Center')

    def test_export_has_a_category_type_column(self):
        import io
        from openpyxl import load_workbook
        from src.map_agent.export import to_xlsx
        actor={'name':'Example University','actor_type':CATEGORIES[0],
               'category_type':'Main Institution','_selected':True}
        workbook=load_workbook(io.BytesIO(to_xlsx({'entities':{'actor':[actor]}})))
        cells=[str(c.value) for sheet in workbook for row in sheet for c in row]
        self.assertIn('Category type',cells)
        self.assertIn('Main Institution',cells)

    def test_seven_regions_cover_every_state_once(self):
        states=[s for codes in REGIONS.values() for s in codes.split()]
        self.assertEqual(len(states),51)
        self.assertEqual(len(set(states)),51)
        for region,codes in REGIONS.items():
            for state in codes.split(): self.assertEqual(region_for('US',state),region)
        self.assertEqual(region_for('USA','New Jersey'),'Mid-Atlantic')
        self.assertEqual(region_for('US','Arizona'),'Mountain West / Rocky Mountain')
        self.assertEqual(region_for('Canada','TX'),'')
        self.assertEqual(region_for('US','NY / NJ'),'')

    def test_new_profile_exports_without_losing_fields(self):
        import io
        from openpyxl import load_workbook
        from pptx import Presentation
        from src.map_agent.export import to_xlsx, to_pptx
        actor={'name':'Example University','actor_type':CATEGORIES[0], '_selected':True,
               'region':'Mid-Atlantic',
               'category_details':[{'label':'Program','value':'Agricultural research ' * 100}],
               'field_evidence':[{'field':'description','value':'Evidence '*200,'source_url':URL,
                   'confidence':'medium','origin':'newly found','date_read':'2026-09-29'}]}
        result={'entities':{'actor':[actor]}}
        workbook=load_workbook(io.BytesIO(to_xlsx(result)))
        cells=[str(c.value) for sheet in workbook for row in sheet for c in row]
        self.assertTrue(any('Agricultural research' in c for c in cells))
        deck=Presentation(io.BytesIO(to_pptx(result)))
        text=' '.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
        self.assertIn('Agricultural research',text)
        self.assertIn('Mid-Atlantic',text)
        self.assertGreater(len(deck.slides),2)
