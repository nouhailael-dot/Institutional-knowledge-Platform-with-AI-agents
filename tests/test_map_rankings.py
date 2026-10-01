"""Offline ranking validation and export coverage; no API calls."""
import unittest
from copy import deepcopy
from src.map_agent.rankings import filter_rankings
from src.map_agent.export import _field_value
from src.map_agent.dedup import deduplicate

URL = "https://www.topuniversities.com/example"
ROW = {"system": "QS World University Rankings", "rank": "201–250", "year": 2026,
       "scope": "overall", "subject": "", "source_url": URL}


class RankingTests(unittest.TestCase):
    def actor(self, **changes):
        return {"name": "Example University", "actor_type": "university", "_entity_type": "actor",
                "rankings": [deepcopy(ROW)], **changes}

    def test_supported_band_year_and_source_preserved(self):
        actor = self.actor()
        filter_rankings(actor, {URL: {}})
        self.assertEqual(actor["rankings"], [ROW])
        self.assertIn("201–250 (2026)", _field_value(actor, "rankings"))
        self.assertIn(URL, _field_value(actor, "rankings"))

    def test_export_survives_a_selection_that_rejected_everything(self):
        import io
        from openpyxl import load_workbook
        from src.map_agent.export import to_xlsx
        rejected = {"entities": {"actor": [{"name": "Example University", "_selected": False}]}}
        cells = [str(c.value) for sheet in load_workbook(io.BytesIO(to_xlsx(rejected)))
                 for row in sheet for c in row]
        self.assertTrue(any("Example University" in c for c in cells))
        empty = load_workbook(io.BytesIO(to_xlsx({"entities": {}})))
        self.assertEqual(len(empty.sheetnames), 1)

    def test_alias_is_dropped_but_a_named_unit_is_not_the_institution(self):
        from src.map_agent.rankings import institution_name, is_ranked_institution, same_institution
        self.assertEqual(institution_name("Massachusetts Institute of Technology (MIT)"),
                         "Massachusetts Institute of Technology")
        unit = "Columbia University (Office of Environmental Stewardship / Sustainable Columbia)"
        self.assertEqual(institution_name(unit), unit)
        self.assertTrue(is_ranked_institution(self.actor(name="Columbia University (CU)")))
        self.assertFalse(is_ranked_institution(self.actor(name=unit)))
        self.assertTrue(same_institution("Columbia  University.", "columbia university"))
        self.assertFalse(same_institution("Columbia University", "Columbia College"))
        self.assertFalse(same_institution("", ""))

    def test_nonuniversity_and_university_lab_do_not_inherit_ranks(self):
        for change in [{"actor_type": "company"}, {"actor_type": "lab"},
                       {"name": "Example University Water Center"}, {"name": "Example University Department of Chemistry"}]:
            actor = self.actor(**change)
            filter_rankings(actor, {URL: {}})
            self.assertNotIn("rankings", actor)
            self.assertEqual(_field_value(actor, "rankings"), "")

    def test_rejects_missing_sources_and_incomplete_or_other_products(self):
        for change in [{"system": "ARWU"}, {"system": "QS Sustainability Rankings"},
                       {"year": None}, {"rank": "top university"}, {"scope": "national"},
                       {"scope": "subject", "subject": ""}, {"source_url": "https://invented.example"}]:
            actor = self.actor(rankings=[{**ROW, **change}])
            filter_rankings(actor, {URL: {}})
            self.assertEqual(actor["rankings"], [])

    def test_subject_scope_and_both_publishers_survive_merge(self):
        second = {**ROW, "system": "Times Higher Education World University Rankings",
                  "rank": "=25", "scope": "subject", "subject": "Engineering"}
        actor = deduplicate([self.actor(), self.actor(rankings=[second])])[0]
        filter_rankings(actor, {URL: {}})
        self.assertEqual(len(actor["rankings"]), 2)
        self.assertIn("Engineering", _field_value(actor, "rankings"))

    def test_duplicates_and_absent_rankings(self):
        actor = self.actor(rankings=[ROW, ROW])
        filter_rankings(actor, {URL: {}})
        self.assertEqual(len(actor["rankings"]), 1)
        actor = self.actor(rankings=None)
        filter_rankings(actor, {URL: {}})
        self.assertEqual(actor["rankings"], [])

class OfficialLookupTests(unittest.TestCase):
    def test_only_official_sources_and_matching_university_are_used(self):
        from unittest.mock import Mock, patch
        from src.map_agent.research import lookup_university_ranking
        actor = {'name': 'Example University', 'actor_type': 'university'}
        sources = Mock()
        sources.search.return_value = [
            {'url': 'https://topuniversities.com.evil.example/rank', 'text': 'bad'},
            {'url': URL, 'text': 'rank evidence'}]
        sources.page.return_value = {'url': URL, 'text': 'rank evidence'}
        rows = [{'name': 'Other University', 'actor_type': 'university', 'rankings': [{**ROW, 'rank': '1'}]},
                {'name': actor['name'], 'actor_type': 'university', 'rankings': [ROW]}]
        with patch('src.map_agent.research.extract', return_value=(rows, '')):
            lookup_university_ranking(object(), actor, ROW['system'], 'topuniversities.com', sources, [])
        sources.page.assert_called_once_with(URL)
        self.assertEqual(actor['rankings'], [ROW])

    def test_missing_evidence_preserves_existing_rank_and_reports_gap(self):
        from unittest.mock import Mock, patch
        from src.map_agent.research import lookup_university_ranking
        actor = {'name': 'Example University', 'rankings': [ROW]}
        sources = Mock()
        sources.search.return_value = []
        gaps = []
        with patch('src.map_agent.research.extract', return_value=([], 'No readable evidence returned.')):
            lookup_university_ranking(object(), actor, ROW['system'], 'topuniversities.com', sources, gaps)
        self.assertEqual(actor['rankings'], [ROW])
        self.assertIn('topuniversities.com', gaps[0])
        sources.page.assert_not_called()
