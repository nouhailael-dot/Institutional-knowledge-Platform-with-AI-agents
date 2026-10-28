"""Export the Phase 1 organizations for a GHUS reviewer to mark Keep or Delete.

    python -m scripts.export_phase1_keep_delete --out Phase1_keep_or_delete.xlsx

Reads the map database named by DATABASE_URL (or --database-url-env) in a
transaction that is declared read only and verified as such before any query;
nothing is written to any database. The workbook lists every actor imported
from the Phase 1 mapping (source_origin = imported_phase1) that still stands on
its own: those merged into another record are counted, not listed. Each row has
a Keep / Delete choice and a comment for the reviewer, and its actor id so the
answers can be applied later through Review.
"""

from __future__ import annotations

import argparse
import os
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from src.db import psycopg_database_url  # noqa: E402
from src.review.taxonomy import REGIONS, category_label, type_label  # noqa: E402

QUERY = """
SELECT a.actor_id::text AS actor_id, a.name, a.status, a.actor_category, a.category_type,
       a.website, a.city, a.state_code, a.region_code, a.country, a.alt_names,
       a.description, a.category_note,
       (SELECT string_agg(DISTINCT s.name, '; ') FROM actor_sector x JOIN sector s
          ON s.sector_id = x.sector_id WHERE x.actor_id = a.actor_id) AS sectors,
       (SELECT string_agg(DISTINCT h.name, '; ') FROM hub_actor ha JOIN hub h
          ON h.hub_id = ha.hub_id WHERE ha.actor_id = a.actor_id) AS hubs
  FROM actor a
 WHERE a.source_origin = 'imported_phase1' AND a.merged_into_actor_id IS NULL
 ORDER BY lower(a.name)
"""
MERGED = ("SELECT count(*) FROM actor WHERE source_origin = 'imported_phase1' "
          "AND merged_into_actor_id IS NOT NULL")

COLUMNS = [  # header, width, key
    ("Decision (Keep / Delete)", 16, None), ("Comment", 34, None),
    ("Organization", 44, "name"), ("Category", 22, "category"), ("Category type", 26, "type"),
    ("Website", 34, "website"), ("City", 16, "city"), ("State", 8, "state_code"),
    ("Region", 26, "region"), ("Country", 14, "country"), ("Sectors", 26, "sectors"),
    ("Hubs", 34, "hubs"), ("Other names", 26, "other_names"), ("Description", 70, "description"),
    ("Category note", 34, "category_note"), ("Status on the map now", 14, "status"),
    ("Actor ID (do not edit)", 38, "actor_id"),
]
MAIN_SHEET = "Phase 1 organizations"
REJECTED_SHEET = "Already rejected"

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="D7410B")      # UM6P orange, as the map's exports
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")       # cells the reviewer fills in
THIN = Side(style="thin", color="D9D4CA")
REGION_NAMES = {code: label for code, label, _states in REGIONS}
# Written by migration 491 on the 58 organizations GHUS confirmed on 2026-10-06;
# GHUS removed it on 2026-10-28 (proposed migration 596). Not shown to a reviewer
# while the live map still holds it.
REMOVED_NOTES = {"Category confirmed by GHUS on 2026-10-06 (Phase 12 proposal)."}


def fetch(database_url: str) -> tuple[list[dict], int, datetime]:
    import psycopg
    from psycopg.rows import dict_row

    with psycopg.connect(psycopg_database_url(database_url), connect_timeout=20,
                         row_factory=dict_row) as conn:
        # A transaction pooler can ignore startup options, so read-only is
        # declared on the transaction itself and checked before any query.
        conn.execute("SET TRANSACTION READ ONLY")
        if conn.execute("SHOW transaction_read_only").fetchone()["transaction_read_only"] != "on":
            raise SystemExit("The database refused read-only mode; nothing was exported.")
        rows = conn.execute(QUERY).fetchall()
        merged = conn.execute(MERGED).fetchone()["count"]
        conn.rollback()
    return [dict(r) for r in rows], int(merged), datetime.now()


def shape(row: dict) -> dict:
    alt = row.get("alt_names")
    return {
        **row,
        "category": category_label(row["actor_category"]) if row.get("actor_category") else "",
        "type": (type_label(row["actor_category"], row["category_type"])
                 if row.get("actor_category") and row.get("category_type") else ""),
        "region": REGION_NAMES.get(row.get("region_code") or "", row.get("region_code") or ""),
        "other_names": "; ".join(alt) if isinstance(alt, list) else (alt or ""),
        "status": str(row.get("status") or "").capitalize(),
        "category_note": None if row.get("category_note") in REMOVED_NOTES else row.get("category_note"),
    }


def write_sheet(ws, rows: list[dict]) -> None:
    for col, (header, width, _key) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 30
    for r, row in enumerate(rows, start=2):
        for col, (_header, _width, key) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=col, value=(row.get(key) or None) if key else None)
            cell.font = FONT
            cell.alignment = Alignment(vertical="top", wrap_text=key in {"description", "category_note"})
            cell.border = Border(bottom=THIN)
            if key is None:
                cell.fill = INPUT_FILL
            if key == "website" and row.get("website"):
                cell.hyperlink = row["website"]
    last = len(rows) + 1
    choice = DataValidation(type="list", formula1='"Keep,Delete"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Keep or Delete",
                            error="Choose Keep or Delete from the list.")
    choice.add(f"A2:A{max(last, 2)}")
    ws.add_data_validation(choice)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLUMNS)).column_letter}{max(last, 2)}"


def write_instructions(ws, *, reviewer: str, main: int, rejected: int, merged: int,
                       exported: datetime) -> None:
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 70
    lines = [
        ("GHUS Phase 1 organizations — Keep or Delete", Font(name="Arial", size=14, bold=True)),
        (f"Prepared for {reviewer}. Exported {exported:%m/%d/%Y %I:%M %p} from the live GHUS map "
         "(read only: this file changes nothing on the map).", FONT),
        ("", FONT),
        ("What to do", BOLD),
        (f"1. In the sheet “{MAIN_SHEET}”, choose Keep or Delete in the yellow Decision column "
         "for every organization (a list appears in the cell).", FONT),
        ("2. Add a comment in the yellow Comment column when it helps: why, or which record it "
         "duplicates.", FONT),
        ("3. Use the filters on the header row to work by category, state or hub.", FONT),
        (f"4. The sheet “{REJECTED_SHEET}” lists the Phase 1 records already rejected. Mark Keep "
         "only for one that should come back; leave the rest empty.", FONT),
        ("", FONT),
        ("What the choices mean", BOLD),
        ("Keep — the organization stays on the map as it is.", FONT),
        ("Delete — the GHUS team will reject it through Review: it is marked Rejected and hidden "
         "from Browse and Ask. Nothing is erased; its history is kept.", FONT),
        ("Nothing happens until the GHUS team applies the decisions; please do not edit the "
         "other columns or the Actor ID.", FONT),
        ("", FONT),
        ("Legend", BOLD),
    ]
    for r, (text, font) in enumerate(lines, start=1):
        cell = ws.cell(row=r, column=1, value=text or None)
        cell.font = font
    legend = len(lines) + 1
    ws.cell(row=legend, column=1, value=None).fill = INPUT_FILL
    ws.cell(row=legend, column=2, value="Yellow cells: fill these in").font = FONT
    ws.cell(row=legend + 1, column=2, value="White cells: information from the map").font = FONT

    example = legend + 3
    ws.cell(row=example, column=1, value="Example of a filled-in row").font = BOLD
    for col, (text, fill) in enumerate([("Delete", True),
                                        ("Duplicate of the Water Council (Milwaukee, WI)", True),
                                        ("The Water Council Inc.", False)], start=1):
        cell = ws.cell(row=example + 1, column=col, value=text)
        cell.font = FONT
        if fill:
            cell.fill = INPUT_FILL

    counts = example + 3
    ws.cell(row=counts, column=1, value="Progress (updates as you fill in)").font = BOLD
    main_ref, rej_ref = f"'{MAIN_SHEET}'", f"'{REJECTED_SHEET}'"
    rows = [
        ("Organizations to review", f"=COUNTA({main_ref}!C2:C{main + 1})"),
        ("Marked Keep", f'=COUNTIF({main_ref}!A2:A{main + 1},"Keep")'),
        ("Marked Delete", f'=COUNTIF({main_ref}!A2:A{main + 1},"Delete")'),
        ("Not decided yet", f"=B{counts + 1}-B{counts + 2}-B{counts + 3}"),
        ("Already rejected (listed separately)", f"=COUNTA({rej_ref}!C2:C{max(rejected + 1, 2)})"),
        ("Of those, marked Keep to bring back", f'=COUNTIF({rej_ref}!A2:A{max(rejected + 1, 2)},"Keep")'),
        ("Merged into another record (not listed)", merged),
    ]
    for offset, (label, value) in enumerate(rows, start=1):
        ws.cell(row=counts + offset, column=1, value=label).font = FONT
        cell = ws.cell(row=counts + offset, column=2, value=value)
        cell.font = FONT
    note = ws.cell(row=counts + len(rows) + 1, column=1,
                   value=f"Merged count as read from the map database on {exported:%m/%d/%Y}.")
    note.font = Font(name="Arial", size=9, italic=True, color="697079")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", required=True, help="the .xlsx file to write")
    parser.add_argument("--reviewer", default="Fatima Zahra")
    parser.add_argument("--database-url-env", default="DATABASE_URL")
    args = parser.parse_args(argv)
    url = os.environ.get(args.database_url_env)
    if not url:
        from dotenv import dotenv_values

        url = dotenv_values(ROOT / ".env").get(args.database_url_env)
    if not url:
        raise SystemExit(f"{args.database_url_env} is not set.")
    rows, merged, exported = fetch(url)
    rows = [shape(r) for r in rows]
    standing = [r for r in rows if r["status"].lower() != "rejected"]
    rejected = [r for r in rows if r["status"].lower() == "rejected"]

    wb = Workbook()
    instructions = wb.active
    instructions.title = "Instructions"
    write_sheet(wb.create_sheet(MAIN_SHEET), standing)
    write_sheet(wb.create_sheet(REJECTED_SHEET), rejected)
    write_instructions(instructions, reviewer=args.reviewer, main=len(standing),
                       rejected=len(rejected), merged=merged, exported=exported)
    wb.calculation.fullCalcOnLoad = True   # the progress counts compute when the file opens
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"{out}: {len(standing)} to review, {len(rejected)} already rejected, "
          f"{merged} merged (not listed)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
